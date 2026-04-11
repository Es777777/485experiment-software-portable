from __future__ import annotations

import argparse
import json
import os
import struct
import sys
import time
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any, Dict, List, Optional, Sequence, Tuple, cast
from zipfile import BadZipFile

from runtime_paths import get_config_runtime_base

if TYPE_CHECKING:
    import serial as serial_module
    from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook
    from openpyxl.worksheet.worksheet import Worksheet

try:
    import serial  # type: ignore
except ImportError:
    serial = None

try:
    from openpyxl import Workbook, load_workbook  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore
except ImportError:
    Workbook = None
    load_workbook = None
    get_column_letter = None


TYPE_REGISTER_COUNT = {
    "uint16": 1,
    "int16": 1,
    "uint32": 2,
    "int32": 2,
    "float32": 2,
}

MODBUS_EXCEPTION_CODES = {
    1: "Illegal function",
    2: "Illegal data address",
    3: "Illegal data value",
    4: "Slave device failure",
    5: "Acknowledge",
    6: "Slave device busy",
    8: "Memory parity error",
    10: "Gateway path unavailable",
    11: "Gateway target failed to respond",
}

REFERENCE_VERSION_REGISTER = 60000
REFERENCE_STARTUP_TARE_ADDRESS_BY_VERSION = {
    0: 21,
    1: 21,
    2: 38,
    3: 38,
    4: 38,
    5: 38,
    99: 21,
}


@dataclass(frozen=True)
class FieldConfig:
    name: str
    function_code: int
    address: int
    data_type: str
    scale: float = 1.0
    offset: float = 0.0
    precision: Optional[int] = None
    byte_order: str = "big"
    word_order: str = "big"

    @property
    def register_count(self) -> int:
        return TYPE_REGISTER_COUNT[self.data_type]


@dataclass(frozen=True)
class AppConfig:
    port: str
    baudrate: int
    bytesize: int
    parity: str
    stopbits: float
    timeout: float
    slave_id: int
    poll_interval_seconds: float
    inter_request_delay_seconds: float
    retries: int
    workbook_path: Path
    workbook_name_timestamp_format: str
    sheet_name: str
    timestamp_format: str
    autosave_every_rows: int
    autosave_interval_seconds: float
    tare_on_startup: bool
    startup_tare_use_reference_mapping: bool
    startup_tare_address: int
    startup_tare_value: int
    startup_tare_settle_seconds: float
    fields: List[FieldConfig]


@dataclass(frozen=True)
class FieldGroup:
    function_code: int
    start_address: int
    quantity: int
    fields: List[FieldConfig]


class ExcelLogger:
    def __init__(
        self,
        workbook_path: Path,
        sheet_name: str,
        field_names: Sequence[str],
        autosave_every_rows: int,
        autosave_interval_seconds: float,
    ) -> None:
        self.workbook_path = workbook_path
        self.sheet_name = sheet_name
        self.autosave_every_rows = autosave_every_rows
        self.autosave_interval_seconds = autosave_interval_seconds
        self.unsaved_rows = 0
        self.pending_rows: List[List[Any]] = []
        self.expected_headers = [
            "timestamp",
            "unix_time",
            "port",
            "baudrate",
            "slave_id",
            "status",
            "error",
            "raw_frames",
        ] + list(field_names)

        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)
        self.headers = self._initialize_workbook()
        self.last_save_monotonic = time.monotonic()
        self.unsaved_rows = 0

    def _initialize_workbook(self) -> List[str]:
        workbook, worksheet = self._load_or_create_workbook()
        try:
            headers = self._ensure_headers(worksheet)
            self._refresh_sheet_layout(worksheet, headers)
            self._save_workbook(workbook)
            return headers
        finally:
            self._close_workbook(workbook)

    def _load_or_create_workbook(self) -> Tuple["OpenpyxlWorkbook", "Worksheet"]:
        assert load_workbook is not None
        assert Workbook is not None
        if self.workbook_path.exists():
            try:
                workbook = cast("OpenpyxlWorkbook", load_workbook(self.workbook_path))
            except Exception as exc:
                self._recover_invalid_workbook(exc)
                workbook = cast("OpenpyxlWorkbook", Workbook())
                worksheet = cast("Worksheet", workbook.active)
                worksheet.title = self.sheet_name
                return workbook, worksheet
            if self.sheet_name in workbook.sheetnames:
                worksheet = cast("Worksheet", workbook[self.sheet_name])
            else:
                worksheet = cast("Worksheet", workbook.create_sheet(self.sheet_name))
            return workbook, worksheet

        workbook = cast("OpenpyxlWorkbook", Workbook())
        worksheet = cast("Worksheet", workbook.active)
        worksheet.title = self.sheet_name
        return workbook, worksheet

    def _build_recovery_path(self, stem_suffix: str, timestamp: str) -> Path:
        candidate = self.workbook_path.with_name(
            "{0}.{1}-{2}{3}".format(
                self.workbook_path.stem,
                stem_suffix,
                timestamp,
                self.workbook_path.suffix,
            )
        )
        counter = 1
        while candidate.exists():
            candidate = self.workbook_path.with_name(
                "{0}.{1}-{2}-{3}{4}".format(
                    self.workbook_path.stem,
                    stem_suffix,
                    timestamp,
                    counter,
                    self.workbook_path.suffix,
                )
            )
            counter += 1
        return candidate

    def _recover_invalid_workbook(self, exc: Exception) -> None:
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup_path = self._build_recovery_path("corrupt", timestamp)
        try:
            self.workbook_path.replace(backup_path)
            print(
                "Warning: {0} could not be opened ({1}). Backed it up to {2}.".format(
                    self.workbook_path,
                    exc,
                    backup_path,
                )
            )
            return
        except PermissionError:
            new_workbook_path = self._build_recovery_path("recovered", timestamp)
            print(
                "Warning: {0} could not be opened ({1}) and is locked by another program. "
                "Logging will continue in {2}.".format(
                    self.workbook_path,
                    exc,
                    new_workbook_path,
                )
            )
            self.workbook_path = new_workbook_path

    def _ensure_headers(self, worksheet: "Worksheet") -> List[str]:
        existing_headers: List[str] = []
        if worksheet.max_row >= 1:
            first_row = [cell.value for cell in worksheet[1]]
            existing_headers = [str(value) for value in first_row if value not in (None, "")]

        if not existing_headers:
            for column_index, header in enumerate(self.expected_headers, start=1):
                worksheet.cell(row=1, column=column_index, value=header)
            return list(self.expected_headers)

        headers = list(existing_headers)
        for header in self.expected_headers:
            if header not in headers:
                worksheet.cell(row=1, column=len(headers) + 1, value=header)
                headers.append(header)
        return headers

    def _refresh_sheet_layout(self, worksheet: "Worksheet", headers: Sequence[str]) -> None:
        assert get_column_letter is not None
        worksheet.freeze_panes = "A2"
        if not headers:
            return

        last_column = get_column_letter(len(headers))
        worksheet.auto_filter.ref = "A1:{0}{1}".format(last_column, max(worksheet.max_row, 1))

        for column_index, header in enumerate(headers, start=1):
            column_letter = get_column_letter(column_index)
            if header == "timestamp":
                width = 22
            elif header == "error":
                width = 40
            elif header == "raw_frames":
                width = 56
            else:
                width = max(14, min(len(header) + 2, 24))
            worksheet.column_dimensions[column_letter].width = width

    def _save_workbook(self, workbook: "OpenpyxlWorkbook") -> None:
        temp_path = self.workbook_path.with_name(
            "{0}.tmp{1}".format(self.workbook_path.stem, self.workbook_path.suffix)
        )
        try:
            workbook.save(temp_path)
            os.replace(temp_path, self.workbook_path)
        except PermissionError as exc:
            if temp_path.exists():
                temp_path.unlink(missing_ok=True)
            raise PermissionError(
                "Cannot write {0}. Close Excel or any other program using the file, then retry."
                .format(self.workbook_path)
            ) from exc
        except Exception:
            if temp_path.exists():
                temp_path.unlink(missing_ok=True)
            raise

    def _close_workbook(self, workbook: "OpenpyxlWorkbook") -> None:
        close_method = getattr(workbook, "close", None)
        if callable(close_method):
            close_method()

    def append_row(self, row_data: Dict[str, Any]) -> None:
        ordered_row = [row_data.get(header, "") for header in self.headers]
        self.pending_rows.append(ordered_row)
        self.unsaved_rows += 1
        if self._should_flush():
            self.flush()

    def _should_flush(self) -> bool:
        if self.unsaved_rows >= self.autosave_every_rows:
            return True
        if self.unsaved_rows > 0 and time.monotonic() - self.last_save_monotonic >= self.autosave_interval_seconds:
            return True
        return False

    def flush(self) -> None:
        if not self.pending_rows and self.workbook_path.exists():
            return

        workbook, worksheet = self._load_or_create_workbook()
        try:
            headers = self._ensure_headers(worksheet)
            if headers != self.headers:
                self.headers = headers
            for row in self.pending_rows:
                worksheet.append(row)
            self._refresh_sheet_layout(worksheet, self.headers)
            self._save_workbook(workbook)
        finally:
            self._close_workbook(workbook)

        self.pending_rows.clear()
        self.unsaved_rows = 0
        self.last_save_monotonic = time.monotonic()

    def close(self) -> None:
        if self.pending_rows:
            self.flush()


def ensure_dependencies() -> None:
    missing = []
    if serial is None:
        missing.append("pyserial")
    if Workbook is None or load_workbook is None or get_column_letter is None:
        missing.append("openpyxl")

    if missing:
        package_list = " ".join(missing)
        raise SystemExit(
            "Missing dependencies: {0}\n"
            "Install them with: py -m pip install -r requirements.txt\n"
            "Or install directly: py -m pip install {1}".format(
                ", ".join(missing), package_list
            )
        )


def load_config(config_path: Path) -> AppConfig:
    config_path = config_path.resolve()
    if not config_path.exists():
        raise FileNotFoundError(
            "Config file not found: {0}. Copy and edit logger_config.json first.".format(config_path)
        )

    payload = json.loads(config_path.read_text(encoding="utf-8"))
    base_dir = get_config_runtime_base(config_path)
    workbook_name_timestamp_format = str(payload.get("workbook_name_timestamp_format", "%Y%m%d_%H%M%S"))
    session_start = datetime.now()
    workbook_path_raw = str(payload.get("workbook_path", "logs/485_data_{start_time}.xlsx"))
    workbook_path = Path(
        workbook_path_raw.replace("{start_time}", session_start.strftime(workbook_name_timestamp_format))
    )
    if not workbook_path.is_absolute():
        workbook_path = base_dir / workbook_path

    fields_payload = payload.get("fields", [])
    if not fields_payload:
        raise ValueError("At least one field must be configured in logger_config.json")

    fields: List[FieldConfig] = []
    for raw_field in fields_payload:
        field = FieldConfig(
            name=str(raw_field["name"]),
            function_code=int(raw_field["function_code"]),
            address=int(raw_field["address"]),
            data_type=str(raw_field["data_type"]).lower(),
            scale=float(raw_field.get("scale", 1.0)),
            offset=float(raw_field.get("offset", 0.0)),
            precision=(
                int(raw_field["precision"])
                if raw_field.get("precision") is not None
                else None
            ),
            byte_order=str(raw_field.get("byte_order", "big")).lower(),
            word_order=str(raw_field.get("word_order", "big")).lower(),
        )
        validate_field(field)
        fields.append(field)

    app_config = AppConfig(
        port=str(payload["port"]),
        baudrate=int(payload.get("baudrate", 9600)),
        bytesize=int(payload.get("bytesize", 8)),
        parity=str(payload.get("parity", "N")).upper(),
        stopbits=float(payload.get("stopbits", 1)),
        timeout=float(payload.get("timeout", 1.0)),
        slave_id=int(payload.get("slave_id", 1)),
        poll_interval_seconds=float(payload.get("poll_interval_seconds", 5)),
        inter_request_delay_seconds=float(payload.get("inter_request_delay_seconds", 0.05)),
        retries=int(payload.get("retries", 1)),
        workbook_path=workbook_path,
        workbook_name_timestamp_format=workbook_name_timestamp_format,
        sheet_name=str(payload.get("sheet_name", "data")),
        timestamp_format=str(payload.get("timestamp_format", "%Y-%m-%d %H:%M:%S")),
        autosave_every_rows=int(payload.get("autosave_every_rows", 10)),
        autosave_interval_seconds=float(payload.get("autosave_interval_seconds", 2.0)),
        tare_on_startup=bool(payload.get("tare_on_startup", True)),
        startup_tare_use_reference_mapping=bool(payload.get("startup_tare_use_reference_mapping", True)),
        startup_tare_address=int(payload.get("startup_tare_address", 21)),
        startup_tare_value=int(payload.get("startup_tare_value", 1)),
        startup_tare_settle_seconds=float(payload.get("startup_tare_settle_seconds", 0.2)),
        fields=fields,
    )
    validate_app_config(app_config)
    return app_config


def validate_field(field: FieldConfig) -> None:
    if field.function_code not in (3, 4):
        raise ValueError("Field {0}: function_code only supports 3 or 4".format(field.name))
    if field.data_type not in TYPE_REGISTER_COUNT:
        raise ValueError(
            "Field {0}: unsupported data_type {1}. Supported: {2}".format(
                field.name, field.data_type, ", ".join(sorted(TYPE_REGISTER_COUNT))
            )
        )
    if field.byte_order not in ("big", "little"):
        raise ValueError("Field {0}: byte_order must be big or little".format(field.name))
    if field.word_order not in ("big", "little"):
        raise ValueError("Field {0}: word_order must be big or little".format(field.name))
    if field.address < 0:
        raise ValueError("Field {0}: address cannot be negative".format(field.name))


def validate_app_config(config: AppConfig) -> None:
    if not 1 <= config.slave_id <= 247:
        raise ValueError("slave_id must be between 1 and 247")
    if config.poll_interval_seconds <= 0:
        raise ValueError("poll_interval_seconds must be greater than 0")
    if config.timeout <= 0:
        raise ValueError("timeout must be greater than 0")
    if config.retries < 0:
        raise ValueError("retries cannot be negative")
    if config.autosave_every_rows <= 0:
        raise ValueError("autosave_every_rows must be greater than 0")
    if config.autosave_interval_seconds <= 0:
        raise ValueError("autosave_interval_seconds must be greater than 0")
    if config.startup_tare_address < 0:
        raise ValueError("startup_tare_address cannot be negative")
    if not 0 <= config.startup_tare_value <= 65535:
        raise ValueError("startup_tare_value must be between 0 and 65535")
    if config.startup_tare_settle_seconds < 0:
        raise ValueError("startup_tare_settle_seconds cannot be negative")
    if config.parity not in ("N", "E", "O", "M", "S"):
        raise ValueError("parity must be one of N, E, O, M, S")
    if config.stopbits not in (1.0, 1.5, 2.0):
        raise ValueError("stopbits must be 1, 1.5, or 2")
    if config.bytesize not in (5, 6, 7, 8):
        raise ValueError("bytesize must be between 5 and 8")


def create_serial_client(config: AppConfig):
    assert serial is not None
    parity_map = {
        "N": serial.PARITY_NONE,
        "E": serial.PARITY_EVEN,
        "O": serial.PARITY_ODD,
        "M": serial.PARITY_MARK,
        "S": serial.PARITY_SPACE,
    }
    stopbits_map = {
        1.0: serial.STOPBITS_ONE,
        1.5: serial.STOPBITS_ONE_POINT_FIVE,
        2.0: serial.STOPBITS_TWO,
    }

    return serial.Serial(
        port=config.port,
        baudrate=config.baudrate,
        bytesize=config.bytesize,
        parity=parity_map[config.parity],
        stopbits=stopbits_map[config.stopbits],
        timeout=config.timeout,
    )


def modbus_crc(frame: bytes) -> int:
    crc = 0xFFFF
    for value in frame:
        crc ^= value
        for _ in range(8):
            if crc & 0x0001:
                crc = (crc >> 1) ^ 0xA001
            else:
                crc >>= 1
    return crc & 0xFFFF


def append_crc(frame: bytes) -> bytes:
    crc = modbus_crc(frame)
    return frame + bytes((crc & 0xFF, (crc >> 8) & 0xFF))


def build_request(slave_id: int, function_code: int, address: int, quantity: int) -> bytes:
    return append_crc(struct.pack(">BBHH", slave_id, function_code, address, quantity))


def build_write_single_request(slave_id: int, address: int, value: int) -> bytes:
    return append_crc(struct.pack(">BBHH", slave_id, 6, address, value))


def validate_response_crc(frame: bytes) -> None:
    if len(frame) < 4:
        raise IOError("Incomplete Modbus response")
    expected_crc = modbus_crc(frame[:-2])
    actual_crc = frame[-2] | (frame[-1] << 8)
    if expected_crc != actual_crc:
        raise IOError(
            "CRC check failed: expected 0x{0:04X}, got 0x{1:04X}".format(expected_crc, actual_crc)
        )


def read_registers(client, slave_id: int, function_code: int, address: int, quantity: int) -> Tuple[List[int], bytes]:
    request = build_request(slave_id, function_code, address, quantity)
    client.reset_input_buffer()
    client.write(request)
    client.flush()

    expected_length = 5 + quantity * 2
    response = client.read(expected_length)
    if len(response) < 5:
        raise TimeoutError(
            "No complete response for address {0}, received {1} bytes".format(address, len(response))
        )

    validate_response_crc(response)

    if response[0] != slave_id:
        raise IOError(
            "Unexpected slave id: expected {0}, got {1}".format(slave_id, response[0])
        )

    if response[1] == (function_code | 0x80):
        exception_code = response[2]
        message = MODBUS_EXCEPTION_CODES.get(exception_code, "Unknown Modbus exception")
        raise IOError(
            "Modbus exception 0x{0:02X}: {1}".format(exception_code, message)
        )

    if response[1] != function_code:
        raise IOError(
            "Unexpected function code: expected {0}, got {1}".format(function_code, response[1])
        )

    byte_count = response[2]
    expected_byte_count = quantity * 2
    if byte_count != expected_byte_count:
        raise IOError(
            "Unexpected byte count: expected {0}, got {1}".format(expected_byte_count, byte_count)
        )

    if len(response) != expected_length:
        raise IOError(
            "Unexpected response length: expected {0}, got {1}".format(expected_length, len(response))
        )

    data_bytes = response[3:-2]
    registers = []
    for index in range(0, len(data_bytes), 2):
        registers.append(int.from_bytes(data_bytes[index:index + 2], byteorder="big", signed=False))
    return registers, response


def write_single_register(client, slave_id: int, address: int, value: int) -> bytes:
    request = build_write_single_request(slave_id, address, value)
    client.reset_input_buffer()
    client.write(request)
    client.flush()

    response = client.read(8)
    if len(response) < 5:
        raise TimeoutError(
            "No complete write response for address {0}, received {1} bytes".format(address, len(response))
        )

    validate_response_crc(response)

    if response[0] != slave_id:
        raise IOError(
            "Unexpected slave id: expected {0}, got {1}".format(slave_id, response[0])
        )

    if response[1] == 0x86:
        exception_code = response[2]
        message = MODBUS_EXCEPTION_CODES.get(exception_code, "Unknown Modbus exception")
        raise IOError(
            "Modbus exception 0x{0:02X}: {1}".format(exception_code, message)
        )

    if response[1] != 6:
        raise IOError(
            "Unexpected function code: expected 6, got {0}".format(response[1])
        )

    if len(response) != 8:
        raise IOError(
            "Unexpected response length: expected 8, got {0}".format(len(response))
        )

    if response[:6] != request[:6]:
        raise IOError(
            "Write confirmation mismatch for address {0}: sent {1}, got {2}".format(
                address,
                request[:6].hex(" ").upper(),
                response[:6].hex(" ").upper(),
            )
        )

    return response


def registers_to_bytes(registers: Sequence[int], byte_order: str, word_order: str) -> bytes:
    chunks = [value.to_bytes(2, byteorder="big", signed=False) for value in registers]
    if byte_order == "little":
        chunks = [chunk[::-1] for chunk in chunks]
    if word_order == "little":
        chunks = list(reversed(chunks))
    return b"".join(chunks)


def decode_value(field: FieldConfig, registers: Sequence[int]) -> Any:
    payload = registers_to_bytes(registers, field.byte_order, field.word_order)

    if field.data_type == "uint16":
        value = int.from_bytes(payload, byteorder="big", signed=False)
    elif field.data_type == "int16":
        value = int.from_bytes(payload, byteorder="big", signed=True)
    elif field.data_type == "uint32":
        value = int.from_bytes(payload, byteorder="big", signed=False)
    elif field.data_type == "int32":
        value = int.from_bytes(payload, byteorder="big", signed=True)
    elif field.data_type == "float32":
        value = struct.unpack(">f", payload)[0]
    else:
        raise ValueError("Unsupported data_type: {0}".format(field.data_type))

    if field.scale != 1.0 or field.offset != 0.0:
        value = value * field.scale + field.offset
    if field.precision is not None:
        value = round(value, field.precision)
    return value


def read_registers_with_retries(
    client,
    config: AppConfig,
    function_code: int,
    address: int,
    quantity: int,
) -> Tuple[List[int], bytes]:
    last_error = None
    total_attempts = config.retries + 1
    for attempt in range(total_attempts):
        try:
            return read_registers(
                client=client,
                slave_id=config.slave_id,
                function_code=function_code,
                address=address,
                quantity=quantity,
            )
        except Exception as exc:
            last_error = exc
            if attempt + 1 < total_attempts:
                time.sleep(0.2)

    raise last_error  # type: ignore[misc]


def write_single_register_with_retries(
    client,
    config: AppConfig,
    address: int,
    value: int,
) -> bytes:
    last_error = None
    total_attempts = config.retries + 1
    for attempt in range(total_attempts):
        try:
            return write_single_register(
                client=client,
                slave_id=config.slave_id,
                address=address,
                value=value,
            )
        except Exception as exc:
            last_error = exc
            if attempt + 1 < total_attempts:
                time.sleep(0.2)

    raise last_error  # type: ignore[misc]


def apply_startup_tare(client, config: AppConfig) -> Tuple[int, bytes]:
    tare_address = config.startup_tare_address
    if config.startup_tare_use_reference_mapping:
        try:
            registers, _ = read_registers_with_retries(
                client=client,
                config=config,
                function_code=3,
                address=REFERENCE_VERSION_REGISTER,
                quantity=1,
            )
            device_version = registers[0]
            tare_address = REFERENCE_STARTUP_TARE_ADDRESS_BY_VERSION.get(device_version, tare_address)
            print(
                "Detected device version {0}; startup tare will use register {1}.".format(
                    device_version,
                    tare_address,
                )
            )
        except Exception as exc:
            print(
                "Warning: could not detect device version from register {0} ({1}); falling back to startup_tare_address={2}.".format(
                    REFERENCE_VERSION_REGISTER,
                    exc,
                    tare_address,
                )
            )

    response = write_single_register_with_retries(
        client=client,
        config=config,
        address=tare_address,
        value=config.startup_tare_value,
    )
    if config.startup_tare_settle_seconds > 0:
        time.sleep(config.startup_tare_settle_seconds)
    return tare_address, response


def build_field_groups(fields: Sequence[FieldConfig]) -> List[FieldGroup]:
    groups: List[FieldGroup] = []
    current_fields: List[FieldConfig] = []
    current_function_code: Optional[int] = None
    start_address = 0
    end_address = 0

    for field in sorted(fields, key=lambda item: (item.function_code, item.address)):
        field_end = field.address + field.register_count
        if not current_fields:
            current_fields = [field]
            current_function_code = field.function_code
            start_address = field.address
            end_address = field_end
            continue

        if field.function_code == current_function_code and field.address <= end_address:
            current_fields.append(field)
            end_address = max(end_address, field_end)
            continue

        groups.append(
            FieldGroup(
                function_code=current_function_code if current_function_code is not None else current_fields[0].function_code,
                start_address=start_address,
                quantity=end_address - start_address,
                fields=list(current_fields),
            )
        )
        current_fields = [field]
        current_function_code = field.function_code
        start_address = field.address
        end_address = field_end

    if current_fields:
        groups.append(
            FieldGroup(
                function_code=current_function_code if current_function_code is not None else current_fields[0].function_code,
                start_address=start_address,
                quantity=end_address - start_address,
                fields=list(current_fields),
            )
        )

    return groups


def poll_once(client, config: AppConfig, logger: ExcelLogger) -> Dict[str, Any]:
    timestamp = datetime.now()
    row_data: Dict[str, Any] = {
        "timestamp": timestamp.strftime(config.timestamp_format),
        "unix_time": round(time.time(), 3),
        "port": config.port,
        "baudrate": config.baudrate,
        "slave_id": config.slave_id,
    }

    raw_frames: Dict[str, str] = {}
    errors: List[str] = []

    for group in build_field_groups(config.fields):
        try:
            registers, response = read_registers_with_retries(
                client=client,
                config=config,
                function_code=group.function_code,
                address=group.start_address,
                quantity=group.quantity,
            )
            frame_hex = response.hex(" ").upper()
            for field in group.fields:
                offset = field.address - group.start_address
                field_registers = registers[offset:offset + field.register_count]
                row_data[field.name] = decode_value(field, field_registers)
                raw_frames[field.name] = frame_hex
        except Exception as exc:
            for field in group.fields:
                row_data[field.name] = ""
                raw_frames[field.name] = ""
                errors.append("{0}: {1}".format(field.name, exc))

        if config.inter_request_delay_seconds > 0:
            time.sleep(config.inter_request_delay_seconds)

    row_data["status"] = "ok" if not errors else "partial_error"
    row_data["error"] = "; ".join(errors)
    row_data["raw_frames"] = json.dumps(raw_frames, ensure_ascii=False)

    logger.append_row(row_data)
    return row_data


def format_console_output(row_data: Dict[str, Any], fields: Sequence[FieldConfig], workbook_path: Path) -> str:
    values = []
    for field in fields:
        value = row_data.get(field.name, "")
        values.append("{0}={1}".format(field.name, value))

    return "[{0}] {1} | {2} | saved to {3}".format(
        row_data["timestamp"],
        row_data["status"],
        ", ".join(values),
        workbook_path,
    )


def run_logger(config: AppConfig, once: bool) -> int:
    logger = ExcelLogger(
        workbook_path=config.workbook_path,
        sheet_name=config.sheet_name,
        field_names=[field.name for field in config.fields],
        autosave_every_rows=config.autosave_every_rows,
        autosave_interval_seconds=config.autosave_interval_seconds,
    )

    try:
        with create_serial_client(config) as client:
            if config.tare_on_startup:
                tare_address, tare_response = apply_startup_tare(client, config)
                print(
                    "Startup tare applied via Modbus 06 at register {0} with value {1} | response {2}".format(
                        tare_address,
                        config.startup_tare_value,
                        tare_response.hex(" ").upper(),
                    )
                )
            while True:
                cycle_started = time.perf_counter()
                row_data = poll_once(client, config, logger)
                print(format_console_output(row_data, config.fields, logger.workbook_path))
                if row_data.get("error"):
                    print("Errors: {0}".format(row_data["error"]))

                if once:
                    return 0

                delay = config.poll_interval_seconds - (time.perf_counter() - cycle_started)
                if delay > 0:
                    time.sleep(delay)
    finally:
        logger.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Read Modbus RTU data from an RS485 transmitter and save it into Excel."
    )
    parser.add_argument(
        "--config",
        default="config/logger_config.json",
        help="Path to the JSON config file. Default: config/logger_config.json",
    )
    parser.add_argument(
        "--port",
        help="Override serial port from config, for example COM4.",
    )
    parser.add_argument(
        "--once",
        action="store_true",
        help="Read one sample only, then exit.",
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    ensure_dependencies()
    parser = build_parser()
    args = parser.parse_args(argv)

    config = load_config(Path(args.config))
    if args.port:
        config = replace(config, port=str(args.port).strip())
    try:
        return run_logger(config, once=args.once)
    except KeyboardInterrupt:
        print("Stopped by user.")
        return 0
    except Exception as exc:
        print("Error: {0}".format(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
