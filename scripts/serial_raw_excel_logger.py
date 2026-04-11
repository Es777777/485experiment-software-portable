from __future__ import annotations

import argparse
import json
import sys
import time
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from runtime_paths import get_config_runtime_base

try:
    import serial  # type: ignore
except ImportError:
    serial = None

try:
    from openpyxl import Workbook, load_workbook  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore
except ImportError:
    Workbook = None
    load_workbook = None
    get_column_letter = None


@dataclass(frozen=True)
class RawAppConfig:
    port: str
    baudrate: int
    bytesize: int
    parity: str
    stopbits: float
    timeout: float
    read_mode: str
    chunk_size: int
    idle_sleep_seconds: float
    workbook_path: Path
    sheet_name: str
    timestamp_format: str
    encoding: str
    decode_errors: str


class RawExcelLogger:
    def __init__(self, workbook_path: Path, sheet_name: str) -> None:
        self.workbook_path = workbook_path
        self.sheet_name = sheet_name
        self.headers = [
            "timestamp",
            "unix_time",
            "port",
            "baudrate",
            "bytesize",
            "parity",
            "stopbits",
            "byte_count",
            "raw_hex",
            "raw_text",
        ]

        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)
        self.workbook, self.worksheet = self._load_or_create_workbook()
        self._ensure_headers()
        self._refresh_sheet_layout()
        self.workbook.save(self.workbook_path)

    def _load_or_create_workbook(self):
        if self.workbook_path.exists():
            workbook = load_workbook(self.workbook_path)
            if self.sheet_name in workbook.sheetnames:
                worksheet = workbook[self.sheet_name]
            else:
                worksheet = workbook.create_sheet(self.sheet_name)
            return workbook, worksheet

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.sheet_name
        return workbook, worksheet

    def _ensure_headers(self) -> None:
        existing_headers: List[str] = []
        if self.worksheet.max_row >= 1:
            first_row = [cell.value for cell in self.worksheet[1]]
            existing_headers = [str(value) for value in first_row if value not in (None, "")]

        if not existing_headers:
            for column_index, header in enumerate(self.headers, start=1):
                self.worksheet.cell(row=1, column=column_index, value=header)
            return

        for header in self.headers:
            if header not in existing_headers:
                self.worksheet.cell(row=1, column=len(existing_headers) + 1, value=header)
                existing_headers.append(header)
        self.headers = existing_headers

    def _refresh_sheet_layout(self) -> None:
        self.worksheet.freeze_panes = "A2"
        last_column = get_column_letter(len(self.headers))
        self.worksheet.auto_filter.ref = "A1:{0}{1}".format(last_column, max(self.worksheet.max_row, 1))

        for column_index, header in enumerate(self.headers, start=1):
            column_letter = get_column_letter(column_index)
            if header == "timestamp":
                width = 22
            elif header == "raw_hex":
                width = 68
            elif header == "raw_text":
                width = 48
            else:
                width = max(14, min(len(header) + 2, 24))
            self.worksheet.column_dimensions[column_letter].width = width

    def append_row(self, row_data: Dict[str, Any]) -> None:
        self.worksheet.append([row_data.get(header, "") for header in self.headers])
        self._refresh_sheet_layout()
        self.workbook.save(self.workbook_path)


def ensure_dependencies() -> None:
    missing = []
    if serial is None:
        missing.append("pyserial")
    if Workbook is None or load_workbook is None or get_column_letter is None:
        missing.append("openpyxl")

    if missing:
        raise SystemExit(
            "Missing dependencies: {0}\n"
            "Install them with: py -m pip install -r requirements.txt".format(
                ", ".join(missing)
            )
        )


def load_config(config_path: Path) -> RawAppConfig:
    config_path = config_path.resolve()
    if not config_path.exists():
        raise FileNotFoundError(
            "Config file not found: {0}. Copy and edit raw_logger_config.json first.".format(config_path)
        )

    payload = json.loads(config_path.read_text(encoding="utf-8"))
    base_dir = get_config_runtime_base(config_path)
    workbook_path = Path(payload.get("workbook_path", "logs/485_raw_data.xlsx"))
    if not workbook_path.is_absolute():
        workbook_path = base_dir / workbook_path

    config = RawAppConfig(
        port=str(payload["port"]),
        baudrate=int(payload.get("baudrate", 9600)),
        bytesize=int(payload.get("bytesize", 8)),
        parity=str(payload.get("parity", "N")).upper(),
        stopbits=float(payload.get("stopbits", 1)),
        timeout=float(payload.get("timeout", 1.0)),
        read_mode=str(payload.get("read_mode", "chunk")).lower(),
        chunk_size=int(payload.get("chunk_size", 256)),
        idle_sleep_seconds=float(payload.get("idle_sleep_seconds", 0.2)),
        workbook_path=workbook_path,
        sheet_name=str(payload.get("sheet_name", "raw_data")),
        timestamp_format=str(payload.get("timestamp_format", "%Y-%m-%d %H:%M:%S")),
        encoding=str(payload.get("encoding", "utf-8")),
        decode_errors=str(payload.get("decode_errors", "replace")),
    )
    validate_config(config)
    return config


def validate_config(config: RawAppConfig) -> None:
    if config.timeout <= 0:
        raise ValueError("timeout must be greater than 0")
    if config.read_mode not in ("chunk", "line"):
        raise ValueError("read_mode must be chunk or line")
    if config.chunk_size <= 0:
        raise ValueError("chunk_size must be greater than 0")
    if config.idle_sleep_seconds < 0:
        raise ValueError("idle_sleep_seconds cannot be negative")
    if config.parity not in ("N", "E", "O", "M", "S"):
        raise ValueError("parity must be one of N, E, O, M, S")
    if config.stopbits not in (1.0, 1.5, 2.0):
        raise ValueError("stopbits must be 1, 1.5, or 2")
    if config.bytesize not in (5, 6, 7, 8):
        raise ValueError("bytesize must be between 5 and 8")


def create_serial_client(config: RawAppConfig):
    parity_map = {
        "N": serial.PARITY_NONE,
        "E": serial.PARITY_EVEN,
        "O": serial.PARITY_ODD,
        "M": serial.PARITY_MARK,
        "S": serial.PARITY_SPACE,
    }
    stopbits_map = {
        1.0: serial.STOPBITS_ONE,
        1.5: serial.STOPBITS_ONE_POINT_FIVE,
        2.0: serial.STOPBITS_TWO,
    }

    return serial.Serial(
        port=config.port,
        baudrate=config.baudrate,
        bytesize=config.bytesize,
        parity=parity_map[config.parity],
        stopbits=stopbits_map[config.stopbits],
        timeout=config.timeout,
    )


def read_payload(client, config: RawAppConfig) -> bytes:
    if config.read_mode == "line":
        return client.readline()

    available = client.in_waiting
    if available > 0:
        return client.read(available)
    return client.read(config.chunk_size)


def build_row(config: RawAppConfig, payload: bytes) -> Dict[str, Any]:
    now = datetime.now()
    return {
        "timestamp": now.strftime(config.timestamp_format),
        "unix_time": round(time.time(), 3),
        "port": config.port,
        "baudrate": config.baudrate,
        "bytesize": config.bytesize,
        "parity": config.parity,
        "stopbits": config.stopbits,
        "byte_count": len(payload),
        "raw_hex": payload.hex(" ").upper(),
        "raw_text": payload.decode(config.encoding, errors=config.decode_errors).strip(),
    }


def format_console_output(row_data: Dict[str, Any], workbook_path: Path) -> str:
    preview = row_data["raw_text"] or row_data["raw_hex"]
    if len(preview) > 60:
        preview = preview[:57] + "..."
    return "[{0}] bytes={1} | {2} | saved to {3}".format(
        row_data["timestamp"],
        row_data["byte_count"],
        preview,
        workbook_path,
    )


def run_logger(config: RawAppConfig, once: bool) -> int:
    logger = RawExcelLogger(config.workbook_path, config.sheet_name)

    with create_serial_client(config) as client:
        print(
            "Listening on {0} ({1},{2},{3},{4}) in {5} mode...".format(
                config.port,
                config.baudrate,
                config.bytesize,
                config.parity,
                config.stopbits,
                config.read_mode,
            )
        )
        while True:
            payload = read_payload(client, config)
            if not payload:
                if config.idle_sleep_seconds > 0:
                    time.sleep(config.idle_sleep_seconds)
                continue

            row_data = build_row(config, payload)
            logger.append_row(row_data)
            print(format_console_output(row_data, config.workbook_path))

            if once:
                return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Passively read serial data from an RS485 transmitter and save it into Excel."
    )
    parser.add_argument(
        "--config",
        default="config/raw_logger_config.json",
        help="Path to the JSON config file. Default: config/raw_logger_config.json",
    )
    parser.add_argument(
        "--port",
        help="Override serial port from config, for example COM4.",
    )
    parser.add_argument(
        "--once",
        action="store_true",
        help="Save one received packet only, then exit.",
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    ensure_dependencies()
    parser = build_parser()
    args = parser.parse_args(argv)

    config = load_config(Path(args.config))
    if args.port:
        config = replace(config, port=str(args.port).strip())
    try:
        return run_logger(config, once=args.once)
    except KeyboardInterrupt:
        print("Stopped by user.")
        return 0
    except Exception as exc:
        print("Error: {0}".format(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
