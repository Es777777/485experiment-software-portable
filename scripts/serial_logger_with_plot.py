"""Real-time serial data logger with live force-time plot.

Reads Modbus RTU data from serial port, writes to Excel, and plots force in real-time.
Single script, no conflicts.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import struct
import sys
import time
import tkinter as tk
import webbrowser
from dataclasses import dataclass, field, replace
from datetime import datetime, timedelta
from pathlib import Path
from tkinter import messagebox, ttk
from typing import Any, Dict, List, Optional, Sequence, Tuple, cast

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import matplotlib

matplotlib.use(os.environ.get("SERIAL_LOGGER_MPL_BACKEND", "TkAgg"))
matplotlib.rcParams["font.sans-serif"] = [
    "Microsoft YaHei UI",
    "Microsoft YaHei",
    "SimHei",
    "SimSun",
    "Arial Unicode MS",
    "DejaVu Sans",
]
matplotlib.rcParams["axes.unicode_minus"] = False
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure
from matplotlib.widgets import Button

try:
    import serial  # type: ignore
except ImportError:
    serial = None

try:
    from openpyxl import Workbook, load_workbook  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore
except ImportError:
    Workbook = None
    load_workbook = None
    get_column_letter = None

try:
    import numpy as np  # type: ignore

    HAS_NUMPY = True
except ImportError:
    HAS_NUMPY = False

from runtime_paths import get_app_root, get_config_runtime_base


TYPE_REGISTER_COUNT = {
    "uint16": 1,
    "int16": 1,
    "uint32": 2,
    "int32": 2,
    "float32": 2,
}

MODBUS_EXCEPTION_CODES = {
    1: "Illegal function",
    2: "Illegal data address",
    3: "Illegal data value",
    4: "Slave device failure",
    5: "Acknowledge",
    6: "Slave device busy",
    8: "Memory parity error",
    10: "Gateway path unavailable",
    11: "Gateway target failed to respond",
}

TOTAL_FORCE_NAME = "total_force"


@dataclass(frozen=True)
class MeasurementSample:
    timestamp: datetime
    raw_values: Dict[str, float]
    total_force: float


@dataclass(frozen=True)
class MeasurementGroupResult:
    index: int
    pwm: str
    start_time: datetime
    end_time: datetime
    total_current: float
    average_total_force: float
    samples: List[MeasurementSample] = field(default_factory=list)
    average_t1: float = 0.0
    average_t2: float = 0.0
    average_t3: float = 0.0
    moment_y: float = 0.0
    moment_x: float = 0.0
    moment_magnitude: float = 0.0
    theta_radians: float = 0.0
    theta_degrees: float = 0.0


@dataclass(frozen=True)
class FormulaDisplaySpec:
    key: str
    expression: str
    value_text: str


@dataclass(frozen=True)
class EditionProfile:
    edition_key: str = "standard"
    display_name: str = "标准版"
    t_channel_scale: float = 1.0
    notes: str = ""


@dataclass
class ActiveMeasurementGroup:
    pwm: str
    start_time: datetime
    samples: List[MeasurementSample]


@dataclass
class MeasurementSessionState:
    completed_groups: List[MeasurementGroupResult] = field(default_factory=list)
    current_group: Optional[ActiveMeasurementGroup] = None


def load_edition_profile(config_dir: Path) -> EditionProfile:
    profile_path = config_dir / "edition_profile.json"
    if not profile_path.exists():
        return EditionProfile()
    payload = json.loads(profile_path.read_text(encoding="utf-8"))
    return EditionProfile(
        edition_key=str(payload.get("edition_key", "standard")),
        display_name=str(payload.get("display_name", "标准版")),
        t_channel_scale=float(payload.get("t_channel_scale", 1.0)),
        notes=str(payload.get("notes", "")),
    )


class SimpleStringVar:
    def __init__(self, value: str = "") -> None:
        self._value = value

    def get(self) -> str:
        return self._value

    def set(self, value: str) -> None:
        self._value = value


@dataclass(frozen=True)
class FieldConfig:
    name: str
    function_code: int
    address: int
    data_type: str
    scale: float = 1.0
    offset: float = 0.0
    precision: Optional[int] = None
    byte_order: str = "big"
    word_order: str = "big"

    @property
    def register_count(self) -> int:
        return TYPE_REGISTER_COUNT[self.data_type]


@dataclass(frozen=True)
class AppConfig:
    port: str
    baudrate: int
    bytesize: int
    parity: str
    stopbits: float
    timeout: float
    slave_id: int
    poll_interval_seconds: float
    inter_request_delay_seconds: float
    retries: int
    workbook_path: Path
    sheet_name: str
    timestamp_format: str
    autosave_every_rows: int
    autosave_interval_seconds: float
    fields: List[FieldConfig]


def configure_text_streams() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is None:
            continue
        try:
            reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def ensure_dependencies() -> None:
    if serial is None:
        raise RuntimeError("缺少 pyserial 依赖，无法打开串口。")
    if Workbook is None or load_workbook is None or get_column_letter is None:
        raise RuntimeError("缺少 openpyxl 依赖，无法写入 Excel。")


def load_config(config_path: Path) -> AppConfig:
    config_path = config_path.resolve()
    payload = json.loads(config_path.read_text(encoding="utf-8"))
    project_root = get_config_runtime_base(config_path)

    workbook_path_str = payload.get("workbook_path", "logs/485_data_{start_time}.xlsx")
    workbook_path = Path(
        workbook_path_str.replace(
            "{start_time}", datetime.now().strftime("%Y%m%d_%H%M%S")
        )
    )
    if not workbook_path.is_absolute():
        workbook_path = project_root / workbook_path

    fields = [
        FieldConfig(
            name=str(f["name"]),
            function_code=int(f["function_code"]),
            address=int(f["address"]),
            data_type=str(f["data_type"]).lower(),
            scale=float(f.get("scale", 1.0)),
            offset=float(f.get("offset", 0.0)),
            precision=int(f["precision"]) if f.get("precision") is not None else None,
            byte_order=str(f.get("byte_order", "big")).lower(),
            word_order=str(f.get("word_order", "big")).lower(),
        )
        for f in payload.get("fields", [])
    ]

    return AppConfig(
        port=str(payload["port"]),
        baudrate=int(payload.get("baudrate", 115200)),
        bytesize=int(payload.get("bytesize", 8)),
        parity=str(payload.get("parity", "N")).upper(),
        stopbits=float(payload.get("stopbits", 1)),
        timeout=float(payload.get("timeout", 1.0)),
        slave_id=int(payload.get("slave_id", 1)),
        poll_interval_seconds=float(payload.get("poll_interval_seconds", 0.1)),
        inter_request_delay_seconds=float(
            payload.get("inter_request_delay_seconds", 0.0)
        ),
        retries=int(payload.get("retries", 1)),
        workbook_path=workbook_path,
        sheet_name=str(payload.get("sheet_name", "data")),
        timestamp_format=str(payload.get("timestamp_format", "%Y-%m-%d %H:%M:%S.%f")),
        autosave_every_rows=int(payload.get("autosave_every_rows", 10)),
        autosave_interval_seconds=float(payload.get("autosave_interval_seconds", 2.0)),
        fields=fields,
    )


def modbus_crc(frame: bytes) -> int:
    crc = 0xFFFF
    for value in frame:
        crc ^= value
        for _ in range(8):
            if crc & 0x0001:
                crc = (crc >> 1) ^ 0xA001
            else:
                crc >>= 1
    return crc & 0xFFFF


def append_crc(frame: bytes) -> bytes:
    crc = modbus_crc(frame)
    return frame + bytes((crc & 0xFF, (crc >> 8) & 0xFF))


def build_request(
    slave_id: int, function_code: int, address: int, quantity: int
) -> bytes:
    return append_crc(struct.pack(">BBHH", slave_id, function_code, address, quantity))


def read_registers(
    client, slave_id: int, function_code: int, address: int, quantity: int
) -> Tuple[List[int], bytes]:
    request = build_request(slave_id, function_code, address, quantity)
    client.reset_input_buffer()
    client.write(request)
    client.flush()
    expected_length = 5 + quantity * 2
    response = client.read(expected_length)
    if len(response) < 5:
        raise TimeoutError("Incomplete response for address {0}".format(address))
    expected_crc = modbus_crc(response[:-2])
    actual_crc = response[-2] | (response[-1] << 8)
    if expected_crc != actual_crc:
        raise IOError("CRC failed")
    if response[1] == (function_code | 0x80):
        raise IOError("Modbus exception 0x{0:02X}".format(response[2]))
    data_bytes = response[3:-2]
    registers = [
        int.from_bytes(data_bytes[i : i + 2], byteorder="big")
        for i in range(0, len(data_bytes), 2)
    ]
    return registers, response


def read_with_retries(
    client, config: AppConfig, function_code: int, address: int, quantity: int
) -> Tuple[List[int], bytes]:
    last_err = None
    for _ in range(config.retries + 1):
        try:
            return read_registers(
                client, config.slave_id, function_code, address, quantity
            )
        except Exception as e:
            last_err = e
            time.sleep(0.2)
    raise last_err  # type: ignore


def registers_to_bytes(
    registers: Sequence[int], byte_order: str, word_order: str
) -> bytes:
    chunks = [v.to_bytes(2, byteorder="big", signed=False) for v in registers]
    if byte_order == "little":
        chunks = [c[::-1] for c in chunks]
    if word_order == "little":
        chunks = list(reversed(chunks))
    return b"".join(chunks)


def decode_value(field: FieldConfig, registers: Sequence[int]) -> Any:
    payload = registers_to_bytes(registers, field.byte_order, field.word_order)
    if field.data_type == "uint16":
        value = int.from_bytes(payload, byteorder="big", signed=False)
    elif field.data_type == "int16":
        value = int.from_bytes(payload, byteorder="big", signed=True)
    elif field.data_type == "uint32":
        value = int.from_bytes(payload, byteorder="big", signed=False)
    elif field.data_type == "int32":
        value = int.from_bytes(payload, byteorder="big", signed=True)
    elif field.data_type == "float32":
        value = struct.unpack(">f", payload)[0]
    else:
        value = int.from_bytes(payload, byteorder="big", signed=False)
    if field.scale != 1.0 or field.offset != 0.0:
        value = value * field.scale + field.offset
    if field.precision is not None:
        value = round(value, field.precision)
    return value


def floor_to_second(ts: datetime) -> datetime:
    return ts.replace(microsecond=0)


def previous_complete_second(ts: datetime) -> datetime:
    return floor_to_second(ts) - timedelta(seconds=1)


def compute_total_force(
    values: Dict[str, float],
    channel_names: Sequence[str],
) -> Optional[float]:
    if any(name not in values for name in channel_names):
        return None
    return sum(values[name] for name in channel_names)


def record_second_bucket(
    buckets: Dict[datetime, Dict[str, List[float]]],
    ts: datetime,
    values: Dict[str, Optional[float]],
) -> None:
    second_key = floor_to_second(ts)
    for name, value in values.items():
        if value is None:
            continue
        bucket = buckets.setdefault(second_key, {})
        bucket.setdefault(name, []).append(value)


def compute_previous_second_averages(
    buckets: Dict[datetime, Dict[str, List[float]]],
    now: datetime,
) -> Tuple[Optional[datetime], Dict[str, float]]:
    second_key = previous_complete_second(now)
    bucket = buckets.get(second_key)
    if not bucket:
        return None, {}

    averages: Dict[str, float] = {}
    for name, samples in bucket.items():
        if samples:
            averages[name] = sum(samples) / len(samples)
    if not averages:
        return None, {}
    return second_key, averages


def build_manual_tare_offsets(
    last_raw_values: Dict[str, float],
    channel_names: Sequence[str],
) -> Dict[str, float]:
    return {
        name: last_raw_values[name] for name in channel_names if name in last_raw_values
    }


def prepare_plot_sample(
    raw_values: Dict[str, float],
    offsets: Dict[str, float],
    channel_names: Sequence[str],
) -> Dict[str, Optional[float]]:
    sample: Dict[str, Optional[float]] = {}
    zeroed_values: Dict[str, float] = {}

    for name in channel_names:
        if name in raw_values:
            zeroed = raw_values[name] - offsets.get(name, 0.0)
            sample[name] = zeroed
            zeroed_values[name] = zeroed
        else:
            sample[name] = None

    sample[TOTAL_FORCE_NAME] = compute_total_force(zeroed_values, channel_names)
    return sample


def append_series_snapshot(
    history: Dict[str, List[float]],
    series_names: Sequence[str],
    sample: Dict[str, Optional[float]],
) -> None:
    for name in series_names:
        value = sample.get(name)
        series = history.setdefault(name, [])
        series.append(float("nan") if value is None else value)


def format_previous_second_title(
    second_ts: Optional[datetime],
    averages: Dict[str, float],
    series_names: Sequence[str],
    tare_message: Optional[str],
) -> str:
    parts: List[str] = []
    if tare_message:
        parts.append(tare_message)

    if second_ts is None or not averages:
        parts.append("等待上一秒统计")
        return " | ".join(parts)

    parts.append("上一秒均值 {0}".format(second_ts.strftime("%H:%M:%S")))
    for name in series_names:
        if name in averages:
            parts.append("{0}={1:.2f}".format(name, averages[name]))
    return " | ".join(parts)


def start_measurement_group(
    state: MeasurementSessionState,
    pwm: str,
    start_time: datetime,
) -> ActiveMeasurementGroup:
    cleaned_pwm = pwm.strip()
    if not cleaned_pwm:
        raise ValueError("请先输入 PWM")
    if state.current_group is not None:
        raise ValueError("当前分组尚未结束，请先点击“结束测量”")
    current_group = ActiveMeasurementGroup(
        pwm=cleaned_pwm, start_time=start_time, samples=[]
    )
    state.current_group = current_group
    return current_group


def record_measurement_sample(
    state: MeasurementSessionState,
    timestamp: datetime,
    raw_values: Dict[str, float],
    total_force: Optional[float],
) -> None:
    if state.current_group is None or total_force is None:
        return
    state.current_group.samples.append(
        MeasurementSample(
            timestamp=timestamp,
            raw_values=dict(raw_values),
            total_force=float(total_force),
        )
    )


def compute_group_channel_averages(
    samples: Sequence[MeasurementSample],
    edition_profile: EditionProfile | None = None,
) -> Tuple[float, float, float]:
    profile = edition_profile or EditionProfile()
    t1_values = [sample.raw_values.get("weight_ch1") for sample in samples]
    t2_values = [sample.raw_values.get("weight_ch2") for sample in samples]
    t3_values = [sample.raw_values.get("weight_ch3") for sample in samples]
    if any(value is None for value in t1_values + t2_values + t3_values):
        raise ValueError("当前分组缺少完整的三路力数据，无法计算力矩")
    return (
        sum(cast(List[float], t1_values)) / len(samples) * profile.t_channel_scale,
        sum(cast(List[float], t2_values)) / len(samples) * profile.t_channel_scale,
        sum(cast(List[float], t3_values)) / len(samples) * profile.t_channel_scale,
    )


def compute_moment_metrics(
    t1: float,
    t2: float,
    t3: float,
) -> Tuple[float, float, float, float, float]:
    moment_y = 0.5 * t1 + 0.5 * t2 - t3
    moment_x = math.sqrt(3.0) / 2.0 * (t2 - t1)
    moment_magnitude = math.sqrt(moment_x**2 + moment_y**2)
    theta_radians = math.atan2(moment_x, moment_y)
    theta_degrees = math.degrees(theta_radians)
    return moment_y, moment_x, moment_magnitude, theta_radians, theta_degrees


def finish_measurement_group(
    state: MeasurementSessionState,
    total_current_text: str,
    end_time: datetime,
    edition_profile: EditionProfile | None = None,
) -> MeasurementGroupResult:
    if state.current_group is None:
        raise ValueError("请先点击“开始测量”")
    cleaned_total_current = total_current_text.strip()
    if not cleaned_total_current:
        raise ValueError("请先输入总电流")
    try:
        total_current = float(cleaned_total_current)
    except ValueError as exc:
        raise ValueError("总电流必须是数字") from exc
    samples = list(state.current_group.samples)
    if not samples:
        raise ValueError("当前分组内还没有有效合力数据，请先采到稳定数据后再结束测量")
    average_t1, average_t2, average_t3 = compute_group_channel_averages(
        samples,
        edition_profile=edition_profile,
    )
    moment_y, moment_x, moment_magnitude, theta_radians, theta_degrees = (
        compute_moment_metrics(average_t1, average_t2, average_t3)
    )
    result = MeasurementGroupResult(
        index=len(state.completed_groups) + 1,
        pwm=state.current_group.pwm,
        start_time=state.current_group.start_time,
        end_time=end_time,
        total_current=total_current,
        average_total_force=sum(sample.total_force for sample in samples)
        / len(samples),
        samples=samples,
        average_t1=average_t1,
        average_t2=average_t2,
        average_t3=average_t3,
        moment_y=moment_y,
        moment_x=moment_x,
        moment_magnitude=moment_magnitude,
        theta_radians=theta_radians,
        theta_degrees=theta_degrees,
    )
    state.completed_groups.append(result)
    state.current_group = None
    return result


def reset_for_next_group(state: MeasurementSessionState) -> None:
    state.current_group = None


def build_measurement_export_path(output_dir: Path, now: datetime) -> Path:
    return output_dir / "measurement_results_{0}.xlsx".format(
        now.strftime("%Y%m%d_%H%M%S")
    )


def build_measurement_export_rows(
    groups: Sequence[MeasurementGroupResult],
) -> Tuple[List[str], List[List[Any]]]:
    max_samples = max((len(group.samples) for group in groups), default=0)
    headers = [
        "序号",
        "PWM",
        "开始时间(北京时间)",
        "结束时间(北京时间)",
        "总电流",
        "合力平均值",
        "T1平均值",
        "T2平均值",
        "T3平均值",
        "M_y",
        "M_x",
        "M",
        "theta(度)",
    ]
    for index in range(1, max_samples + 1):
        headers.extend(
            [
                "原始时间戳_{0}".format(index),
                "原始合力_{0}".format(index),
                "原始ch1_{0}".format(index),
                "原始ch2_{0}".format(index),
                "原始ch3_{0}".format(index),
            ]
        )

    rows: List[List[Any]] = []
    for group in groups:
        row: List[Any] = [
            group.index,
            group.pwm,
            group.start_time.strftime("%Y-%m-%d %H:%M:%S.%f"),
            group.end_time.strftime("%Y-%m-%d %H:%M:%S.%f"),
            group.total_current,
            group.average_total_force,
            group.average_t1,
            group.average_t2,
            group.average_t3,
            group.moment_y,
            group.moment_x,
            group.moment_magnitude,
            group.theta_degrees,
        ]
        for sample in group.samples:
            row.extend(
                [
                    sample.timestamp.strftime("%Y-%m-%d %H:%M:%S.%f"),
                    sample.total_force,
                    sample.raw_values.get("weight_ch1", ""),
                    sample.raw_values.get("weight_ch2", ""),
                    sample.raw_values.get("weight_ch3", ""),
                ]
            )
        for _ in range(len(group.samples), max_samples):
            row.extend(["", "", "", "", ""])
        rows.append(row)
    return headers, rows


def build_moment_formula_specs(
    result: MeasurementGroupResult,
) -> List[FormulaDisplaySpec]:
    return [
        FormulaDisplaySpec(
            key="M_y",
            expression=r"$M_y = \frac{1}{2}T_1 + \frac{1}{2}T_2 - T_3$",
            value_text="{0:.3f}".format(result.moment_y),
        ),
        FormulaDisplaySpec(
            key="M_x",
            expression=r"$M_x = \frac{\sqrt{3}}{2}(T_2 - T_1)$",
            value_text="{0:.3f}".format(result.moment_x),
        ),
        FormulaDisplaySpec(
            key="M",
            expression=r"$M = \sqrt{M_x^2 + M_y^2}$",
            value_text="{0:.3f}".format(result.moment_magnitude),
        ),
        FormulaDisplaySpec(
            key="theta",
            expression=r"$\theta = \operatorname{atan2}(M_x, M_y)$",
            value_text="{0:.2f}°".format(result.theta_degrees),
        ),
    ]


def build_placeholder_formula_specs() -> List[FormulaDisplaySpec]:
    return [
        FormulaDisplaySpec(
            key="M_y",
            expression=r"$M_y = \frac{1}{2}T_1 + \frac{1}{2}T_2 - T_3$",
            value_text="--",
        ),
        FormulaDisplaySpec(
            key="M_x",
            expression=r"$M_x = \frac{\sqrt{3}}{2}(T_2 - T_1)$",
            value_text="--",
        ),
        FormulaDisplaySpec(
            key="M",
            expression=r"$M = \sqrt{M_x^2 + M_y^2}$",
            value_text="--",
        ),
        FormulaDisplaySpec(
            key="theta",
            expression=r"$\theta = \operatorname{atan2}(M_x, M_y)$",
            value_text="--",
        ),
    ]


def build_measurement_metric_values(
    result: Optional[MeasurementGroupResult],
) -> Dict[str, str]:
    if result is None:
        return {
            "average_total_force": "--",
            "T1": "--",
            "T2": "--",
            "T3": "--",
            "total_current": "--",
        }
    return {
        "average_total_force": "{0:.3f}".format(result.average_total_force),
        "T1": "{0:.3f}".format(result.average_t1),
        "T2": "{0:.3f}".format(result.average_t2),
        "T3": "{0:.3f}".format(result.average_t3),
        "total_current": "{0:.3f} A".format(result.total_current),
    }


def build_measurement_context_text(
    completed_count: int,
    status: str,
    current_pwm: str = "",
    latest_result: Optional[MeasurementGroupResult] = None,
    export_name: str = "",
) -> str:
    if export_name:
        return "已完成组数={0} | 已导出 {1} | 可继续检查结果或开始新一轮测量".format(
            completed_count,
            export_name,
        )
    if status == "测量中":
        return "当前组 PWM={0} | 正在采样，等待稳定后结束测量并录入总电流".format(
            current_pwm or "--"
        )
    if latest_result is not None and status == "已完成":
        return (
            "最近完成 #{0} | T1/T2/T3={1:.3f}/{2:.3f}/{3:.3f} | 可继续下一组或导出"
        ).format(
            latest_result.index,
            latest_result.average_t1,
            latest_result.average_t2,
            latest_result.average_t3,
        )
    return "已完成组数={0} | 请输入下一组 PWM，或检查历史后导出".format(completed_count)


def build_measurement_footer_text(profile: EditionProfile | None = None) -> str:
    active = profile or EditionProfile()
    suffix = ""
    if active.edition_key == "buaa":
        suffix = "  |  北航特供版：T1/T2/T3 与由其推导的力矩结果按 1/3 换算"
    return (
        "485 Experiment Software  |  "
        "Maintained by Chenghang Li  |  "
        "github.com/Es777777/485experiment-software-portable" + suffix
    )


def build_measurement_layout_metrics() -> Dict[str, Any]:
    return {
        "hint_wrap": 720,
        "context_wrap": 720,
        "footer_wrap": 720,
        "formula_figure_size": (6.4, 2.6),
    }


def split_measurement_actions(labels: Sequence[str]) -> List[List[str]]:
    items = list(labels)
    midpoint = max(1, len(items) // 2)
    return [items[:midpoint], items[midpoint:]]


def build_measurement_theme_palette() -> Dict[str, str]:
    return {
        "panel_background": "#f4f7fb",
        "card_background": "#ffffff",
        "accent": "#0f4c81",
        "status_active": "#d97706",
        "status_done": "#0f766e",
        "text_primary": "#102a43",
        "text_muted": "#52606d",
        "border": "#d9e2ec",
        "history_background": "#edf2f7",
        "button_primary": "#0f4c81",
    }


def write_measurement_workbook(
    output_path: Path, groups: Sequence[MeasurementGroupResult]
) -> None:
    assert Workbook is not None
    headers, rows = build_measurement_export_rows(groups)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "measurements"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


class ExcelLogger:
    def __init__(
        self,
        workbook_path: Path,
        sheet_name: str,
        field_names: Sequence[str],
        autosave_every_rows: int,
        autosave_interval_seconds: float,
    ) -> None:
        self.workbook_path = workbook_path
        self.sheet_name = sheet_name
        self.autosave_every_rows = autosave_every_rows
        self.autosave_interval_seconds = autosave_interval_seconds
        self.pending_rows: List[List[Any]] = []
        self.last_appended_timestamp: Optional[str] = None
        self.last_flushed_timestamp: Optional[str] = None
        self.expected_headers = [
            "timestamp",
            "unix_time",
            "port",
            "baudrate",
            "slave_id",
            "status",
            "error",
            "raw_frames",
        ] + list(field_names)
        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)
        self.headers = self._initialize_workbook()
        self.last_save_monotonic = time.monotonic()
        self.unsaved_rows = 0

    def _initialize_workbook(self) -> List[str]:
        assert Workbook is not None
        if self.workbook_path.exists():
            try:
                assert load_workbook is not None
                wb = load_workbook(self.workbook_path)
            except Exception:
                wb = Workbook()
        else:
            wb = Workbook()
        ws = wb.active
        if ws is None or ws.title != self.sheet_name:
            ws = wb.create_sheet(self.sheet_name)
        existing = [str(c.value) for c in ws[1] if c.value] if ws.max_row >= 1 else []
        headers = list(existing)
        for h in self.expected_headers:
            if h not in headers:
                ws.cell(row=1, column=len(headers) + 1, value=h)
                headers.append(h)
        self.headers = headers
        self._save(wb)
        wb.close()
        return headers

    def _save(self, wb) -> None:
        ws = wb.active
        if ws is not None:
            ws.freeze_panes = "A2"
            last_col = (
                get_column_letter(len(self.headers)) if get_column_letter else "H"
            )
            ws.auto_filter.ref = "A1:{0}{1}".format(last_col, max(ws.max_row, 1))
        tmp = self.workbook_path.with_name(
            "{0}.tmp{1}".format(self.workbook_path.stem, self.workbook_path.suffix)
        )
        try:
            wb.save(tmp)
            import os as _os

            _os.replace(tmp, self.workbook_path)
        except Exception:
            if tmp.exists():
                tmp.unlink(missing_ok=True)
            raise

    def _timestamp_column_index(self) -> int:
        return self.headers.index("timestamp") + 1

    def _read_last_saved_timestamp(self) -> Optional[str]:
        assert load_workbook is not None
        if not self.workbook_path.exists():
            return None
        wb = load_workbook(self.workbook_path, read_only=True, data_only=True)
        try:
            ws = wb[self.sheet_name] if self.sheet_name in wb.sheetnames else wb.active
            timestamp_col = self._timestamp_column_index()
            for row_index in range(ws.max_row, 1, -1):
                value = ws.cell(row=row_index, column=timestamp_col).value
                if value in (None, ""):
                    continue
                return str(value)
        finally:
            close_method = getattr(wb, "close", None)
            if callable(close_method):
                close_method()
        return None

    def _verify_last_saved_timestamp(self, expected_timestamp: Optional[str]) -> None:
        if not expected_timestamp:
            return
        actual_timestamp = self._read_last_saved_timestamp()
        if actual_timestamp != expected_timestamp:
            raise RuntimeError(
                "Excel 写入校验失败：期望最后时间戳为 {0}，实际为 {1}。"
                " 请检查 Excel 文件是否被占用或保存失败。".format(
                    expected_timestamp,
                    actual_timestamp if actual_timestamp is not None else "空",
                )
            )

    def append_row(self, row_data: Dict[str, Any]) -> None:
        ordered = [row_data.get(h, "") for h in self.headers]
        self.pending_rows.append(ordered)
        timestamp_value = row_data.get("timestamp")
        self.last_appended_timestamp = (
            str(timestamp_value) if timestamp_value not in (None, "") else None
        )
        self.unsaved_rows += 1
        if (
            self.unsaved_rows >= self.autosave_every_rows
            or time.monotonic() - self.last_save_monotonic
            >= self.autosave_interval_seconds
        ):
            self.flush()

    def flush(self) -> None:
        if not self.pending_rows:
            return
        assert load_workbook is not None
        rows_to_save = list(self.pending_rows)
        expected_timestamp = self.last_appended_timestamp
        try:
            wb = load_workbook(self.workbook_path)
            ws = wb[self.sheet_name] if self.sheet_name in wb.sheetnames else wb.active
            for row in rows_to_save:
                ws.append(row)
            self._save(wb)
            wb.close()
            self._verify_last_saved_timestamp(expected_timestamp)
            print(
                "  [Saved {0} rows to {1}]".format(
                    len(rows_to_save), self.workbook_path.name
                )
            )
        except Exception as e:
            print("  [Save error: {0}]".format(e))
            raise
        self.pending_rows = self.pending_rows[len(rows_to_save) :]
        self.unsaved_rows = len(self.pending_rows)
        self.last_flushed_timestamp = expected_timestamp
        self.last_save_monotonic = time.monotonic()

    def close(self) -> None:
        if self.pending_rows:
            self.flush()
        self._verify_last_saved_timestamp(self.last_appended_timestamp)


class LivePlotter:
    COLORS = ["b", "r", "g", "m", "c", "y"]

    def __init__(
        self,
        channel_names: List[str],
        window_seconds: int = 60,
        smooth: bool = False,
        calc_window: int = 10,
        tare_seconds: float = 2.0,
        edition_profile: EditionProfile | None = None,
    ) -> None:
        self.channel_names = channel_names
        self.edition_profile = edition_profile or EditionProfile()
        self.total_force_name = TOTAL_FORCE_NAME
        self.plot_series_names = list(channel_names) + [self.total_force_name]
        self.window_seconds = window_seconds
        self.calc_window = calc_window
        self.smooth = smooth
        self.tare_seconds = tare_seconds
        self.offsets: Dict[str, float] = {name: 0.0 for name in channel_names}
        self.tare_done = False
        self.tare_start: Optional[datetime] = None
        self.tare_buffer: Dict[str, List[float]] = {name: [] for name in channel_names}
        self.last_raw_values: Dict[str, float] = {}
        self.last_tare_message: Optional[str] = None
        self.last_tare_message_until: float = 0.0
        self.second_buckets: Dict[datetime, Dict[str, List[float]]] = {}
        self.last_sample_ts: Optional[datetime] = None
        self.measurement_state = MeasurementSessionState()
        self.measurement_output_dir = get_app_root() / "output"
        self.all_times: List[datetime] = []
        self.all_values: Dict[str, List[float]] = {
            name: [] for name in self.plot_series_names
        }
        self.raw_lines: Dict[str, Any] = {}
        self.smooth_lines: Dict[str, Any] = {}
        self.avg_lines: Dict[str, Any] = {}
        self.measurement_pwm_entry: Any = None
        self.measurement_current_entry: Any = None

        self.fig, self.ax = plt.subplots(figsize=(12, 6))
        plt.subplots_adjust(top=0.86)

        for i, name in enumerate(self.plot_series_names):
            color = self.COLORS[i % len(self.COLORS)]
            (raw_line,) = self.ax.plot(
                [], [], color=color, linewidth=0.8, alpha=0.3, markersize=2
            )
            self.raw_lines[name] = raw_line
            (smooth_line,) = self.ax.plot(
                [], [], color=color, linewidth=2.0, label=name
            )
            self.smooth_lines[name] = smooth_line
            (avg_line,) = self.ax.plot(
                [],
                [],
                color=color,
                linestyle="--",
                linewidth=1.2,
                label="Avg {0}".format(name),
            )
            self.avg_lines[name] = avg_line

        self.button_ax = self.fig.add_axes([0.84, 0.9, 0.12, 0.06])
        self.tare_button = Button(self.button_ax, "去皮")
        self.tare_button.on_clicked(self.request_manual_tare)

        self.ax.set_title("Taring... ({0:.1f}s remaining)".format(tare_seconds))
        self.ax.set_xlabel("Time")
        self.ax.set_ylabel("Force (g)")
        self.ax.grid(True, alpha=0.3)
        self.ax.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M:%S"))
        self.ax.legend(loc="upper right", fontsize=8)
        self.fig.autofmt_xdate()
        plt.ion()
        if "agg" not in matplotlib.get_backend().lower():
            plt.show(block=False)
        self.measurement_var_master = None
        self.measurement_status_var = SimpleStringVar("未开始")
        self.measurement_pwm_var = SimpleStringVar("")
        self.measurement_current_var = SimpleStringVar("")
        self.measurement_summary_var = SimpleStringVar("已完成组数: 0")
        self.measurement_history_var = SimpleStringVar("暂无已完成组")
        self.measurement_context_var = SimpleStringVar(
            build_measurement_context_text(0, "未开始")
        )
        self.measurement_formula_value_vars: Dict[str, Any] = {
            "M_y": SimpleStringVar("--"),
            "M_x": SimpleStringVar("--"),
            "M": SimpleStringVar("--"),
            "theta": SimpleStringVar("--"),
        }
        self.measurement_metric_value_vars: Dict[str, Any] = {
            "average_total_force": SimpleStringVar("--"),
            "T1": SimpleStringVar("--"),
            "T2": SimpleStringVar("--"),
            "T3": SimpleStringVar("--"),
            "total_current": SimpleStringVar("--"),
        }
        self.measurement_footer_var = SimpleStringVar(
            build_measurement_footer_text(self.edition_profile)
        )
        self.measurement_panel: Any = None
        self.measurement_status_label: Any = None
        self.measurement_summary_label: Any = None
        self.measurement_history_label: Any = None
        self.measurement_footer_label: Any = None
        self.measurement_metric_labels: Dict[str, Any] = {}
        self.measurement_formula_figure: Optional[Figure] = None
        self.measurement_formula_canvas: Any = None
        self.measurement_formula_axes: List[Any] = []
        self._build_measurement_panel()

    def _resolve_measurement_var_master(self) -> Any:
        manager = getattr(self.fig.canvas, "manager", None)
        window = getattr(manager, "window", None)
        if isinstance(window, tk.Misc):
            return window
        return None

    def _make_string_var(self, value: str) -> Any:
        if self.measurement_var_master is not None:
            return tk.StringVar(master=self.measurement_var_master, value=value)
        return SimpleStringVar(value)

    def _rebind_measurement_vars(self) -> None:
        self.measurement_status_var = self._make_string_var(
            self.measurement_status_var.get()
        )
        self.measurement_pwm_var = self._make_string_var(self.measurement_pwm_var.get())
        self.measurement_current_var = self._make_string_var(
            self.measurement_current_var.get()
        )
        self.measurement_summary_var = self._make_string_var(
            self.measurement_summary_var.get()
        )
        self.measurement_history_var = self._make_string_var(
            self.measurement_history_var.get()
        )
        self.measurement_context_var = self._make_string_var(
            self.measurement_context_var.get()
        )
        self.measurement_footer_var = self._make_string_var(
            self.measurement_footer_var.get()
        )
        self.measurement_formula_value_vars = {
            key: self._make_string_var(var.get())
            for key, var in self.measurement_formula_value_vars.items()
        }
        self.measurement_metric_value_vars = {
            key: self._make_string_var(var.get())
            for key, var in self.measurement_metric_value_vars.items()
        }

    def _build_measurement_panel(self) -> None:
        manager = getattr(self.fig.canvas, "manager", None)
        window = getattr(manager, "window", None)
        if window is None:
            return
        self.measurement_var_master = self._resolve_measurement_var_master()
        self._rebind_measurement_vars()

        palette = build_measurement_theme_palette()
        metrics = build_measurement_layout_metrics()
        panel = tk.LabelFrame(
            window,
            text="分组测量",
            bg=palette["panel_background"],
            fg=palette["text_primary"],
            padx=10,
            pady=10,
            bd=1,
            relief="groove",
        )
        panel.pack(side="bottom", fill="x", padx=10, pady=8)
        self.measurement_panel = panel
        panel.grid_columnconfigure(0, weight=1)

        header_row = tk.Frame(panel, bg=palette["panel_background"])
        header_row.grid(row=0, column=0, sticky="ew")
        header_row.grid_columnconfigure(0, weight=1)
        input_box = tk.LabelFrame(
            header_row,
            text="输入区",
            bg=palette["panel_background"],
            fg=palette["text_primary"],
            padx=10,
            pady=8,
            bd=1,
            relief="groove",
        )
        input_box.grid(row=0, column=0, sticky="ew")
        input_box.grid_columnconfigure(1, weight=1)
        input_box.grid_columnconfigure(3, weight=1)
        action_box = tk.LabelFrame(
            panel,
            text="操作区",
            bg=palette["panel_background"],
            fg=palette["text_primary"],
            padx=10,
            pady=8,
            bd=1,
            relief="groove",
        )
        action_box.grid(row=1, column=0, sticky="ew", pady=(8, 0))

        tk.Label(
            input_box,
            text="PWM",
            bg=palette["panel_background"],
            fg=palette["text_primary"],
        ).grid(row=0, column=0, padx=(0, 6), pady=4, sticky="w")
        self.measurement_pwm_entry = ttk.Entry(
            input_box,
            textvariable=self.measurement_pwm_var,
            width=12,
        )
        self.measurement_pwm_entry.grid(row=0, column=1, pady=4, sticky="ew")
        tk.Label(
            input_box,
            text="总电流",
            bg=palette["panel_background"],
            fg=palette["text_primary"],
        ).grid(row=0, column=2, padx=(12, 6), pady=4, sticky="w")
        self.measurement_current_entry = ttk.Entry(
            input_box,
            textvariable=self.measurement_current_var,
            width=12,
        )
        self.measurement_current_entry.grid(row=0, column=3, pady=4, sticky="ew")

        action_specs = [
            ("开始测量", self.start_measurement),
            ("结束测量", self.finish_measurement),
            ("下一组", self.prepare_next_group),
            ("导出表格", self.export_measurements),
        ]
        split_labels = split_measurement_actions([label for label, _ in action_specs])
        action_lookup = {label: command for label, command in action_specs}
        for row_index, labels in enumerate(split_labels):
            for column_index, label in enumerate(labels):
                ttk.Button(
                    action_box,
                    text=label,
                    command=action_lookup[label],
                ).grid(
                    row=row_index,
                    column=column_index,
                    sticky="ew",
                    padx=(0 if column_index == 0 else 8, 0),
                    pady=(0 if row_index == 0 else 8, 0),
                )
                action_box.grid_columnconfigure(column_index, weight=1)

        result_box = tk.Frame(
            panel,
            bg=palette["card_background"],
            bd=1,
            relief="solid",
            highlightbackground=palette["border"],
            highlightthickness=1,
        )
        result_box.grid(row=2, column=0, sticky="ew", pady=(10, 0))
        result_box.grid_columnconfigure(0, weight=1)
        result_box.grid_columnconfigure(1, weight=1)

        self.measurement_status_label = tk.Label(
            result_box,
            textvariable=self.measurement_status_var,
            bg=palette["card_background"],
            fg=palette["text_primary"],
            font=("Microsoft YaHei UI", 11, "bold"),
            anchor="w",
        )
        self.measurement_status_label.grid(
            row=0,
            column=0,
            sticky="w",
            padx=12,
            pady=(10, 4),
        )
        self.measurement_summary_label = tk.Label(
            result_box,
            textvariable=self.measurement_summary_var,
            bg=palette["card_background"],
            fg=palette["accent"],
            font=("Microsoft YaHei UI", 10, "bold"),
            anchor="w",
        )
        self.measurement_summary_label.grid(
            row=0,
            column=1,
            sticky="w",
            padx=(12, 12),
            pady=(10, 4),
        )

        hint_label = tk.Label(
            result_box,
            text="操作顺序：输入 PWM -> 点开始测量 -> 稳定后输入总电流 -> 点结束测量 -> 点下一组",
            bg=palette["card_background"],
            fg=palette["text_muted"],
            anchor="w",
            justify="left",
            wraplength=metrics["hint_wrap"],
        )
        hint_label.grid(row=1, column=0, columnspan=2, sticky="ew", padx=12)

        context_label = tk.Label(
            result_box,
            textvariable=self.measurement_context_var,
            bg=palette["card_background"],
            fg=palette["text_primary"],
            anchor="w",
            justify="left",
            font=("Microsoft YaHei UI", 9),
            wraplength=metrics["context_wrap"],
        )
        context_label.grid(
            row=2, column=0, columnspan=2, sticky="ew", padx=12, pady=(6, 0)
        )

        metric_row = tk.Frame(result_box, bg=palette["card_background"])
        metric_row.grid(
            row=3, column=0, columnspan=2, sticky="ew", padx=12, pady=(8, 4)
        )
        metric_labels = [
            ("average_total_force", "平均合力"),
            ("T1", "T1"),
            ("T2", "T2"),
            ("T3", "T3"),
            ("total_current", "总电流"),
        ]
        for column_index, (key, title) in enumerate(metric_labels):
            card = tk.Frame(
                metric_row,
                bg=palette["panel_background"],
                bd=1,
                relief="solid",
                highlightbackground=palette["border"],
                highlightthickness=1,
                padx=10,
                pady=6,
            )
            card.grid(
                row=0,
                column=column_index,
                sticky="nsew",
                padx=(0 if column_index == 0 else 6, 0),
            )
            metric_row.grid_columnconfigure(column_index, weight=1)
            tk.Label(
                card,
                text=title,
                bg=palette["panel_background"],
                fg=palette["text_muted"],
                font=("Microsoft YaHei UI", 9),
                anchor="w",
            ).pack(anchor="w")
            value_label = tk.Label(
                card,
                textvariable=self.measurement_metric_value_vars[key],
                bg=palette["panel_background"],
                fg=palette["accent"],
                font=("Consolas", 12, "bold"),
                anchor="w",
            )
            value_label.pack(anchor="w", pady=(4, 0))
            self.measurement_metric_labels[key] = value_label

        self.measurement_formula_figure = Figure(
            figsize=metrics["formula_figure_size"],
            dpi=100,
            facecolor=palette["card_background"],
        )
        self.measurement_formula_axes = [
            self.measurement_formula_figure.add_subplot(2, 2, index + 1)
            for index in range(4)
        ]
        for axis in self.measurement_formula_axes:
            axis.set_axis_off()
            axis.set_facecolor(palette["card_background"])
        self.measurement_formula_canvas = FigureCanvasTkAgg(
            self.measurement_formula_figure,
            master=result_box,
        )
        formula_widget = self.measurement_formula_canvas.get_tk_widget()
        formula_widget.configure(
            background=palette["card_background"],
            highlightbackground=palette["border"],
            highlightthickness=1,
        )
        formula_widget.grid(row=4, column=0, columnspan=2, sticky="ew", padx=12, pady=8)

        history_box = tk.Frame(
            panel,
            bg=palette["history_background"],
            bd=1,
            relief="solid",
            highlightbackground=palette["border"],
            highlightthickness=1,
        )
        history_box.grid(row=3, column=0, sticky="ew", pady=(10, 0))
        history_box.grid_columnconfigure(0, weight=1)
        tk.Label(
            history_box,
            text="已完成组历史",
            bg=palette["history_background"],
            fg=palette["text_primary"],
            font=("Microsoft YaHei UI", 10, "bold"),
            anchor="w",
        ).grid(row=0, column=0, sticky="w", padx=12, pady=(8, 4))
        self.measurement_history_label = tk.Label(
            history_box,
            textvariable=self.measurement_history_var,
            bg=palette["history_background"],
            fg=palette["text_primary"],
            justify="left",
            anchor="w",
        )
        self.measurement_history_label.grid(
            row=1,
            column=0,
            sticky="ew",
            padx=12,
            pady=(0, 8),
        )

        footer_box = tk.Frame(panel, bg=palette["panel_background"])
        footer_box.grid(row=4, column=0, sticky="ew", pady=(8, 0))
        self.measurement_footer_label = tk.Label(
            footer_box,
            textvariable=self.measurement_footer_var,
            bg=palette["panel_background"],
            fg=palette["text_muted"],
            cursor="hand2",
            anchor="w",
            justify="left",
            wraplength=metrics["footer_wrap"],
        )
        self.measurement_footer_label.grid(row=0, column=0, sticky="w")
        self.measurement_footer_label.bind("<Button-1>", self._open_measurement_github)

        self._apply_measurement_result(None)
        self._update_measurement_visual_state()

    def _open_measurement_github(self, _event: object | None = None) -> None:
        try:
            webbrowser.open(
                "https://github.com/Es777777/485experiment-software-portable"
            )
        except Exception:
            pass

    def _render_formula_specs(self, specs: Sequence[FormulaDisplaySpec]) -> None:
        if self.measurement_formula_figure is None or not self.measurement_formula_axes:
            return
        palette = build_measurement_theme_palette()
        for axis in self.measurement_formula_axes:
            axis.clear()
            axis.set_axis_off()
            axis.set_facecolor(palette["card_background"])
        for axis, spec in zip(self.measurement_formula_axes, specs):
            axis.text(
                0.03,
                0.68,
                spec.expression,
                fontsize=13,
                color=palette["text_primary"],
            )
            axis.text(
                0.05,
                0.2,
                "= {0}".format(spec.value_text),
                fontsize=11,
                color=palette["accent"],
                fontweight="bold",
            )
        self.measurement_formula_figure.tight_layout(pad=1.1)
        if self.measurement_formula_canvas is not None:
            self.measurement_formula_canvas.draw_idle()

    def _apply_measurement_result(
        self,
        result: Optional[MeasurementGroupResult],
    ) -> None:
        metric_values = build_measurement_metric_values(result)
        for key, value in metric_values.items():
            self.measurement_metric_value_vars[key].set(value)
        specs = (
            build_placeholder_formula_specs()
            if result is None
            else build_moment_formula_specs(result)
        )
        for spec in specs:
            self.measurement_formula_value_vars[spec.key].set(spec.value_text)
        self._render_formula_specs(specs)

    def _set_measurement_context(
        self,
        latest_result: Optional[MeasurementGroupResult] = None,
        export_name: str = "",
    ) -> None:
        self.measurement_context_var.set(
            build_measurement_context_text(
                len(self.measurement_state.completed_groups),
                str(self.measurement_status_var.get()),
                self._read_entry_text(
                    self.measurement_pwm_entry,
                    self.measurement_pwm_var,
                ).strip(),
                latest_result=latest_result,
                export_name=export_name,
            )
        )

    def _update_measurement_visual_state(self) -> None:
        palette = build_measurement_theme_palette()
        if self.measurement_status_label is None:
            return
        status = str(self.measurement_status_var.get())
        color = palette["text_primary"]
        if status == "测量中":
            color = palette["status_active"]
        elif status in {"已完成", "已导出"}:
            color = palette["status_done"]
        self.measurement_status_label.configure(fg=color)
        if self.measurement_summary_label is not None:
            self.measurement_summary_label.configure(fg=palette["accent"])
        if self.measurement_footer_label is not None:
            self.measurement_footer_label.configure(fg=palette["text_muted"])

    def _read_entry_text(self, entry: Any, fallback_var: Any) -> str:
        if entry is not None:
            try:
                value = entry.get()
                if value is not None:
                    text = str(value)
                    fallback_var.set(text)
                    return text
            except Exception:
                pass
        return str(fallback_var.get())

    def _refresh_measurement_history(self) -> None:
        if not self.measurement_state.completed_groups:
            self.measurement_history_var.set("暂无已完成组")
            return

        lines = []
        for group in self.measurement_state.completed_groups:
            lines.append(
                "#{0} PWM={1} 开始={2} 结束={3} 总电流={4:.3f} 平均合力={5:.3f} M={6:.3f} θ={7:.2f}°".format(
                    group.index,
                    group.pwm,
                    group.start_time.strftime("%H:%M:%S"),
                    group.end_time.strftime("%H:%M:%S"),
                    group.total_current,
                    group.average_total_force,
                    group.moment_magnitude,
                    group.theta_degrees,
                )
            )
        self.measurement_history_var.set("\n".join(lines))

    def _set_tare_message(self, message: str) -> None:
        self.last_tare_message = message
        self.last_tare_message_until = time.monotonic() + 2.0

    def _active_tare_message(self) -> Optional[str]:
        if time.monotonic() <= self.last_tare_message_until:
            return self.last_tare_message
        return None

    def request_manual_tare(self, _event: object | None = None) -> None:
        if not self.last_raw_values:
            self._set_tare_message("去皮失败：暂无有效数据")
            return

        if self.last_sample_ts is not None:
            self.second_buckets.pop(floor_to_second(self.last_sample_ts), None)

        self.offsets.update(
            build_manual_tare_offsets(self.last_raw_values, self.channel_names)
        )
        self.tare_done = True
        self.tare_start = None
        self.tare_buffer = {name: [] for name in self.channel_names}
        self._set_tare_message("已去皮 {0}".format(datetime.now().strftime("%H:%M:%S")))

    def _tare(self, ts: datetime, values: Dict[str, float]) -> bool:
        if self.tare_start is None:
            self.tare_start = ts
        elapsed = (ts - self.tare_start).total_seconds()
        for name in self.channel_names:
            if name in values:
                self.tare_buffer[name].append(values[name])
        remaining = max(0.0, self.tare_seconds - elapsed)
        self.ax.set_title("Taring... ({0:.1f}s remaining)".format(remaining))
        self.fig.canvas.draw()
        self.fig.canvas.flush_events()
        if elapsed >= self.tare_seconds:
            for name in self.channel_names:
                buf = self.tare_buffer.get(name, [])
                if buf:
                    self.offsets[name] = sum(buf) / len(buf)
            self.tare_done = True
            self._set_tare_message("启动去皮完成")
            return True
        return False

    def _prune_history(self, ts: datetime) -> None:
        cutoff = ts.timestamp() - self.window_seconds
        while self.all_times and self.all_times[0].timestamp() < cutoff:
            self.all_times.pop(0)
            for name in self.plot_series_names:
                if self.all_values[name]:
                    self.all_values[name].pop(0)

        oldest_bucket = floor_to_second(ts) - timedelta(seconds=self.window_seconds + 2)
        for bucket_ts in list(self.second_buckets):
            if bucket_ts < oldest_bucket:
                del self.second_buckets[bucket_ts]

    def add_point(self, ts: datetime, values: Dict[str, float]) -> None:
        self.last_raw_values = dict(values)
        self.last_sample_ts = ts

        if not self.tare_done:
            self._tare(ts, values)
            return

        sample = prepare_plot_sample(values, self.offsets, self.channel_names)
        record_measurement_sample(
            self.measurement_state,
            ts,
            dict(values),
            sample.get(self.total_force_name),
        )

        self.all_times.append(ts)
        append_series_snapshot(self.all_values, self.plot_series_names, sample)
        record_second_bucket(self.second_buckets, ts, sample)
        self._prune_history(ts)
        self._update()

    def _build_smoothed_values(self, values: List[float]) -> List[float]:
        if not HAS_NUMPY or len(values) <= 3 or any(math.isnan(v) for v in values):
            return values
        win = max(3, len(values) // 30)
        kernel = np.ones(win) / win
        padded = np.pad(values, (win // 2, win // 2), mode="edge")
        return list(np.convolve(padded, kernel, mode="valid")[: len(values)])

    def _update(self) -> None:
        if not self.all_times:
            return

        now = self.all_times[-1]
        second_ts, averages = compute_previous_second_averages(self.second_buckets, now)

        for name in self.plot_series_names:
            vals = self.all_values.get(name, [])
            if not vals:
                continue

            self.raw_lines[name].set_data(self.all_times, vals)

            if self.smooth:
                smoothed = self._build_smoothed_values(vals)
                self.smooth_lines[name].set_data(self.all_times, smoothed)
            else:
                self.smooth_lines[name].set_data(self.all_times, vals)

            if name in averages:
                avg_vals = [averages[name]] * len(self.all_times)
                self.avg_lines[name].set_data(self.all_times, avg_vals)
            else:
                self.avg_lines[name].set_data([], [])

        self.ax.relim()
        self.ax.autoscale_view()
        if len(self.all_times) == 1 or self.all_times[0] == self.all_times[-1]:
            single_time = self.all_times[-1]
            self.ax.set_xlim(
                [
                    single_time - timedelta(milliseconds=500),
                    single_time + timedelta(milliseconds=500),
                ]
            )
        else:
            self.ax.set_xlim([self.all_times[0], self.all_times[-1]])
        self.ax.set_title(
            format_previous_second_title(
                second_ts,
                averages,
                self.plot_series_names,
                self._active_tare_message(),
            )
        )
        self.fig.canvas.draw()
        self.fig.canvas.flush_events()

    def start_measurement(self) -> None:
        try:
            start_measurement_group(
                self.measurement_state,
                self._read_entry_text(
                    self.measurement_pwm_entry, self.measurement_pwm_var
                ),
                datetime.now(),
            )
        except ValueError as exc:
            messagebox.showerror("开始测量失败", str(exc))
            return
        self.measurement_status_var.set("测量中")
        self.measurement_summary_var.set(
            "已完成组数: {0} | 当前 PWM={1}".format(
                len(self.measurement_state.completed_groups),
                self._read_entry_text(
                    self.measurement_pwm_entry, self.measurement_pwm_var
                ).strip(),
            )
        )
        self._apply_measurement_result(None)
        self._set_measurement_context()
        self._update_measurement_visual_state()

    def finish_measurement(self) -> None:
        try:
            result = finish_measurement_group(
                self.measurement_state,
                self._read_entry_text(
                    self.measurement_current_entry, self.measurement_current_var
                ),
                datetime.now(),
                edition_profile=self.edition_profile,
            )
        except ValueError as exc:
            messagebox.showerror("结束测量失败", str(exc))
            return
        self.measurement_status_var.set("已完成")
        self.measurement_summary_var.set(
            "已完成组数: {0} | 最近一组 PWM={1} | 平均合力={2:.3f}".format(
                len(self.measurement_state.completed_groups),
                result.pwm,
                result.average_total_force,
            )
        )
        self._apply_measurement_result(result)
        self._set_measurement_context(result)
        self._refresh_measurement_history()
        self._update_measurement_visual_state()

    def prepare_next_group(self) -> None:
        if self.measurement_state.current_group is not None:
            messagebox.showerror("下一组失败", "当前组尚未结束")
            return
        reset_for_next_group(self.measurement_state)
        self.measurement_pwm_var.set("")
        self.measurement_current_var.set("")
        self.measurement_status_var.set("未开始")
        self.measurement_summary_var.set(
            "已完成组数: {0}".format(len(self.measurement_state.completed_groups))
        )
        self._apply_measurement_result(None)
        self._set_measurement_context()
        self._refresh_measurement_history()
        self._update_measurement_visual_state()

    def export_measurements(self) -> Path:
        if not self.measurement_state.completed_groups:
            messagebox.showerror("导出失败", "没有可导出的测量组")
            raise ValueError("No completed measurement groups")
        output_path = build_measurement_export_path(
            self.measurement_output_dir, datetime.now()
        )
        write_measurement_workbook(output_path, self.measurement_state.completed_groups)
        self.measurement_status_var.set("已导出")
        self.measurement_summary_var.set("已导出: {0}".format(output_path.name))
        self._set_measurement_context(export_name=output_path.name)
        self._refresh_measurement_history()
        self._update_measurement_visual_state()
        return output_path

    def close(self) -> None:
        plt.close(self.fig)


def poll_once(client, config: AppConfig, logger: ExcelLogger) -> Dict[str, Any]:
    timestamp = datetime.now()
    row_data: Dict[str, Any] = {
        "timestamp": timestamp.strftime(config.timestamp_format),
        "unix_time": round(time.time(), 3),
        "port": config.port,
        "baudrate": config.baudrate,
        "slave_id": config.slave_id,
    }
    raw_frames: Dict[str, str] = {}
    errors: List[str] = []

    # Group fields by function_code and address range
    sorted_fields = sorted(config.fields, key=lambda f: (f.function_code, f.address))
    groups: List[Tuple[int, int, int, List[FieldConfig]]] = []
    if sorted_fields:
        cur_fc = sorted_fields[0].function_code
        cur_start = sorted_fields[0].address
        cur_end = cur_start + sorted_fields[0].register_count
        cur_fields = [sorted_fields[0]]
        for f in sorted_fields[1:]:
            if f.function_code == cur_fc and f.address <= cur_end:
                cur_fields.append(f)
                cur_end = max(cur_end, f.address + f.register_count)
            else:
                groups.append((cur_fc, cur_start, cur_end - cur_start, cur_fields))
                cur_fc = f.function_code
                cur_start = f.address
                cur_end = cur_start + f.register_count
                cur_fields = [f]
        groups.append((cur_fc, cur_start, cur_end - cur_start, cur_fields))

    for fc, start_addr, qty, grp_fields in groups:
        try:
            registers, response = read_with_retries(client, config, fc, start_addr, qty)
            frame_hex = response.hex(" ").upper()
            for f in grp_fields:
                offset = f.address - start_addr
                field_regs = registers[offset : offset + f.register_count]
                row_data[f.name] = decode_value(f, field_regs)
                raw_frames[f.name] = frame_hex
        except Exception as exc:
            for f in grp_fields:
                row_data[f.name] = ""
                raw_frames[f.name] = ""
                errors.append("{0}: {1}".format(f.name, exc))
        if config.inter_request_delay_seconds > 0:
            time.sleep(config.inter_request_delay_seconds)

    row_data["status"] = "ok" if not errors else "partial_error"
    row_data["error"] = "; ".join(errors)
    row_data["raw_frames"] = json.dumps(raw_frames, ensure_ascii=False)
    logger.append_row(row_data)
    return row_data


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Serial data logger with live force-time plot."
    )
    parser.add_argument(
        "--config", default="config/logger_config.json", help="Config file path."
    )
    parser.add_argument(
        "--port", help="Override serial port from config, for example COM4."
    )
    parser.add_argument(
        "--plot-window", type=int, default=60, help="Plot time window in seconds."
    )
    parser.add_argument(
        "--smooth", action="store_true", help="Apply smoothing to plot."
    )
    parser.add_argument("--once", action="store_true", help="Read one sample and exit.")
    return parser


def resolve_config_path(raw_path: str) -> Path:
    path = Path(raw_path)
    if path.is_absolute():
        return path.resolve()
    return (get_app_root() / path).resolve()


def create_serial_client(config: AppConfig):
    assert serial is not None
    parity_map = {
        "N": serial.PARITY_NONE,
        "E": serial.PARITY_EVEN,
        "O": serial.PARITY_ODD,
        "M": serial.PARITY_MARK,
        "S": serial.PARITY_SPACE,
    }
    stopbits_map = {
        1.0: serial.STOPBITS_ONE,
        1.5: serial.STOPBITS_ONE_POINT_FIVE,
        2.0: serial.STOPBITS_TWO,
    }
    return serial.Serial(
        port=config.port,
        baudrate=config.baudrate,
        bytesize=config.bytesize,
        parity=parity_map[config.parity],
        stopbits=stopbits_map[config.stopbits],
        timeout=config.timeout,
    )


def format_serial_open_error(port: str, exc: BaseException) -> str:
    return (
        "无法打开串口 {0}: {1}\n"
        "请确认设备已连接、电源已打开、驱动已安装，并在主窗口点击“搜索串口”选择实际的 COM 口。"
    ).format(port, exc)


def main(argv: Optional[Sequence[str]] = None) -> int:
    configure_text_streams()
    parser = build_parser()
    args = parser.parse_args(argv)

    try:
        ensure_dependencies()
        config_path = resolve_config_path(args.config)
        config = load_config(config_path)
        edition_profile = load_edition_profile(get_config_runtime_base(config_path))
        if args.port:
            config = replace(config, port=str(args.port).strip())
    except Exception as exc:
        print("实时曲线初始化失败: {0}".format(exc), file=sys.stderr)
        return 1

    print("正在连接串口 {0}...".format(config.port))
    print("保存到: {0}".format(config.workbook_path))
    try:
        client = create_serial_client(config)
    except Exception as exc:
        print(format_serial_open_error(config.port, exc), file=sys.stderr)
        return 1
    print("Connected.")

    logger: ExcelLogger | None = None
    plotter: LivePlotter | None = None

    try:
        logger = ExcelLogger(
            workbook_path=config.workbook_path,
            sheet_name=config.sheet_name,
            field_names=[f.name for f in config.fields],
            autosave_every_rows=config.autosave_every_rows,
            autosave_interval_seconds=config.autosave_interval_seconds,
        )

        plotter = LivePlotter(
            channel_names=[f.name for f in config.fields],
            window_seconds=args.plot_window,
            smooth=args.smooth,
            edition_profile=edition_profile,
        )

        while True:
            cycle_start = time.perf_counter()
            row_data = poll_once(client, config, logger)

            plot_values = {}
            for f in config.fields:
                val = row_data.get(f.name)
                if val is not None and val != "":
                    plot_values[f.name] = float(val)
            if plot_values:
                plotter.add_point(datetime.now(), plot_values)

            print(
                "[{0}] {1} | {2}".format(
                    row_data["timestamp"],
                    row_data["status"],
                    ", ".join(
                        "{0}={1}".format(f.name, row_data.get(f.name, ""))
                        for f in config.fields
                    ),
                )
            )

            if row_data.get("error"):
                print("  Errors: {0}".format(row_data["error"]))

            if args.once:
                break

            delay = config.poll_interval_seconds - (time.perf_counter() - cycle_start)
            if delay > 0:
                time.sleep(delay)

    except KeyboardInterrupt:
        print("\nStopped by user.")
    except Exception as exc:
        print("实时曲线运行失败: {0}".format(exc), file=sys.stderr)
        return 1
    finally:
        if logger is not None:
            logger.close()
        if plotter is not None:
            plotter.close()
        client.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
