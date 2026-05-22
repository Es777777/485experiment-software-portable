from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import cv2  # type: ignore
import numpy as np
from openpyxl import load_workbook  # type: ignore
from rapidocr_onnxruntime import RapidOCR  # type: ignore

from runtime_paths import get_app_root


VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".wmv", ".flv", ".m4v"}
TIMESTAMP_RE = re.compile(r"(20\d{6})[_-]?(\d{6})")

# These ratios are calibrated from the captured power-supply display.
BLUE_CROP = (0.307, 0.042, 0.754, 0.487)
RED_CROP = (0.251, 0.403, 0.754, 0.805)
OUTPUT_HEADERS = [
    "video_timestamp",
    "video_match_delta_ms",
    "video_voltage",
    "video_current",
    "video_file",
]
DIGIT_SEGMENT_PATTERNS = {
    "0": np.array([1, 1, 1, 0, 1, 1, 1], dtype=np.float32),
    "1": np.array([0, 0, 1, 0, 0, 1, 0], dtype=np.float32),
    "2": np.array([1, 0, 1, 1, 1, 0, 1], dtype=np.float32),
    "3": np.array([1, 0, 1, 1, 0, 1, 1], dtype=np.float32),
    "4": np.array([0, 1, 1, 1, 0, 1, 0], dtype=np.float32),
    "5": np.array([1, 1, 0, 1, 0, 1, 1], dtype=np.float32),
    "6": np.array([1, 1, 0, 1, 1, 1, 1], dtype=np.float32),
    "7": np.array([1, 0, 1, 0, 0, 1, 0], dtype=np.float32),
    "8": np.array([1, 1, 1, 1, 1, 1, 1], dtype=np.float32),
    "9": np.array([1, 1, 1, 1, 0, 1, 1], dtype=np.float32),
}
SEGMENT_SAMPLE_BOXES = (
    (0.28, 0.04, 0.72, 0.18),
    (0.06, 0.16, 0.28, 0.46),
    (0.72, 0.16, 0.94, 0.46),
    (0.26, 0.42, 0.74, 0.58),
    (0.06, 0.54, 0.28, 0.84),
    (0.72, 0.54, 0.94, 0.84),
    (0.28, 0.82, 0.72, 0.96),
)
SEGMENT_WEIGHTS = np.array([1.2, 1.0, 1.0, 1.4, 1.0, 1.0, 1.2], dtype=np.float32)


@dataclass(frozen=True)
class VideoInfo:
    path: Path
    start_time: datetime
    fps: float
    frame_count: int
    duration_seconds: float
    roi: Tuple[int, int, int, int]

    @property
    def end_time(self) -> datetime:
        return self.start_time + timedelta(seconds=self.duration_seconds)


class VideoReader:
    def __init__(self, info: VideoInfo) -> None:
        self.info = info
        self.capture = cv2.VideoCapture(str(info.path))
        if not self.capture.isOpened():
            raise RuntimeError("Could not open video: {0}".format(info.path))
        self.cache: Dict[int, Tuple[Optional[float], Optional[float], datetime, float]] = {}

    def close(self) -> None:
        self.capture.release()

    def read_measurement(
        self,
        frame_index: int,
        ocr: RapidOCR,
    ) -> Tuple[Optional[float], Optional[float], datetime, float]:
        frame_index = max(0, min(frame_index, max(self.info.frame_count - 1, 0)))
        cached = self.cache.get(frame_index)
        if cached is not None:
            return cached

        self.capture.set(cv2.CAP_PROP_POS_FRAMES, frame_index)
        ok, frame = self.capture.read()
        if not ok or frame is None:
            raise RuntimeError(
                "Could not read frame {0} from {1}".format(frame_index, self.info.path)
            )

        rotated = rotate_frame_if_needed(frame)
        display = crop_roi(rotated, self.info.roi)
        voltage = extract_blue_value(display, ocr)
        current = extract_red_value(display, ocr)
        frame_seconds = frame_index / self.info.fps if self.info.fps > 0 else 0.0
        frame_time = self.info.start_time + timedelta(seconds=frame_seconds)
        result = (voltage, current, frame_time, frame_seconds)
        self.cache[frame_index] = result
        return result


def rotate_frame_if_needed(frame: np.ndarray) -> np.ndarray:
    if frame.shape[0] > frame.shape[1]:
        return cv2.rotate(frame, cv2.ROTATE_90_CLOCKWISE)
    return frame


def crop_roi(image: np.ndarray, roi: Tuple[int, int, int, int]) -> np.ndarray:
    x1, y1, x2, y2 = roi
    return image[y1:y2, x1:x2]


def crop_by_ratio(image: np.ndarray, ratios: Tuple[float, float, float, float]) -> np.ndarray:
    h, w = image.shape[:2]
    x1 = int(round(w * ratios[0]))
    y1 = int(round(h * ratios[1]))
    x2 = int(round(w * ratios[2]))
    y2 = int(round(h * ratios[3]))
    return image[y1:y2, x1:x2]


def clean_digits(result: Optional[Sequence[Sequence[object]]]) -> Optional[str]:
    if not result:
        return None
    text = "".join(str(item[1]) for item in result if len(item) >= 2)
    digits = "".join(char for char in text if char.isdigit())
    return digits or None


def parse_fixed_one_decimal(digits: Optional[str]) -> Optional[float]:
    if not digits:
        return None
    if len(digits) == 1:
        return float("0.{0}".format(digits))
    return float("{0}.{1}".format(digits[:-1], digits[-1]))


def run_ocr_digits(image: np.ndarray, ocr: RapidOCR) -> Optional[str]:
    result = ocr(image)
    if not result:
        return None
    return clean_digits(result[0])


def threshold_display_digits(crop: np.ndarray, color: str) -> np.ndarray:
    b, g, r = cv2.split(crop)
    if color == "blue":
        primary = b
        secondary = np.maximum(g, r)
        min_level = 60
    else:
        primary = r
        secondary = np.maximum(g, b)
        min_level = 75

    diff = cv2.subtract(primary, secondary)
    diff = cv2.GaussianBlur(diff, (5, 5), 0)
    otsu_level, _ = cv2.threshold(diff, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    level = max(min_level, int(otsu_level))
    _, mask = cv2.threshold(diff, level, 255, cv2.THRESH_BINARY)
    kernel = np.ones((3, 3), dtype=np.uint8)
    mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, kernel, iterations=1)
    mask = cv2.morphologyEx(mask, cv2.MORPH_OPEN, kernel, iterations=1)
    return mask


def build_digit_mask_variants(crop: np.ndarray, color: str) -> List[np.ndarray]:
    base = threshold_display_digits(crop, color)
    variants: List[np.ndarray] = []
    for scale in (4, 5):
        enlarged = cv2.resize(base, None, fx=scale, fy=scale, interpolation=cv2.INTER_NEAREST)
        variants.append(enlarged)
        variants.append(cv2.erode(enlarged, np.ones((2, 2), dtype=np.uint8), iterations=1))
        variants.append(cv2.erode(enlarged, np.ones((3, 3), dtype=np.uint8), iterations=1))
        variants.append(cv2.erode(enlarged, np.ones((3, 1), dtype=np.uint8), iterations=1))
        variants.append(cv2.erode(enlarged, np.ones((2, 2), dtype=np.uint8), iterations=2))
    return variants


def merge_component_boxes(
    left: Tuple[int, int, int, int, int],
    right: Tuple[int, int, int, int, int],
) -> Tuple[int, int, int, int, int]:
    return (
        min(left[0], right[0]),
        min(left[1], right[1]),
        max(left[2], right[2]),
        max(left[3], right[3]),
        left[4] + right[4],
    )


def locate_digit_boxes(mask: np.ndarray, expected_digits: int = 3) -> Optional[List[Tuple[int, int, int, int]]]:
    min_area = max(12, int(mask.size * 0.0005))
    num_labels, _, stats, _ = cv2.connectedComponentsWithStats(mask)
    boxes: List[Tuple[int, int, int, int, int]] = []
    for index in range(1, num_labels):
        x = int(stats[index, cv2.CC_STAT_LEFT])
        y = int(stats[index, cv2.CC_STAT_TOP])
        w = int(stats[index, cv2.CC_STAT_WIDTH])
        h = int(stats[index, cv2.CC_STAT_HEIGHT])
        area = int(stats[index, cv2.CC_STAT_AREA])
        if area < min_area:
            continue
        boxes.append((x, y, x + w, y + h, area))

    boxes.sort(key=lambda item: (item[0], item[1]))
    while len(boxes) > expected_digits:
        median_area = float(np.median([box[4] for box in boxes])) if boxes else 0.0
        merge_index = next(
            (
                index
                for index, box in enumerate(boxes)
                if box[4] <= median_area * 0.45 or (box[2] - box[0]) <= mask.shape[1] * 0.08
            ),
            None,
        )
        if merge_index is None:
            return None

        candidates = []
        if merge_index > 0:
            candidates.append(merge_index - 1)
        if merge_index + 1 < len(boxes):
            candidates.append(merge_index + 1)
        if not candidates:
            return None

        center_x = (boxes[merge_index][0] + boxes[merge_index][2]) / 2.0
        target_index = min(
            candidates,
            key=lambda index: abs(((boxes[index][0] + boxes[index][2]) / 2.0) - center_x),
        )
        left_index = min(merge_index, target_index)
        right_index = max(merge_index, target_index)
        merged = merge_component_boxes(boxes[left_index], boxes[right_index])
        boxes = boxes[:left_index] + [merged] + boxes[left_index + 1 : right_index] + boxes[right_index + 1 :]

    if len(boxes) != expected_digits:
        return None

    heights = [box[3] - box[1] for box in boxes]
    if not heights or min(heights) <= mask.shape[0] * 0.25:
        return None

    return [(box[0], box[1], box[2], box[3]) for box in boxes]


def sample_region(mask: np.ndarray, box: Tuple[float, float, float, float]) -> float:
    h, w = mask.shape[:2]
    x1 = max(0, min(w - 1, int(round(w * box[0]))))
    y1 = max(0, min(h - 1, int(round(h * box[1]))))
    x2 = max(x1 + 1, min(w, int(round(w * box[2]))))
    y2 = max(y1 + 1, min(h, int(round(h * box[3]))))
    region = mask[y1:y2, x1:x2]
    if region.size == 0:
        return 0.0
    return float(region.mean())


def decode_digit_mask(mask: np.ndarray) -> Tuple[Optional[str], float]:
    normalized = cv2.copyMakeBorder(mask, 6, 6, 6, 6, cv2.BORDER_CONSTANT, value=0)
    normalized = cv2.resize(normalized, (72, 128), interpolation=cv2.INTER_NEAREST)
    normalized = (normalized > 0).astype(np.float32)
    segment_levels = np.array([sample_region(normalized, box) for box in SEGMENT_SAMPLE_BOXES], dtype=np.float32)

    best_digit: Optional[str] = None
    best_score = -1.0
    second_score = -1.0
    for digit, pattern in DIGIT_SEGMENT_PATTERNS.items():
        matches = np.where(pattern > 0.5, segment_levels, 1.0 - segment_levels)
        score = float(np.dot(matches, SEGMENT_WEIGHTS) / SEGMENT_WEIGHTS.sum())
        if score > best_score:
            second_score = best_score
            best_digit = digit
            best_score = score
        elif score > second_score:
            second_score = score

    if best_digit is None or best_score < 0.56 or (best_score - second_score) < 0.02:
        return None, best_score
    return best_digit, best_score


def read_seven_segment_digits(mask: np.ndarray, expected_digits: int = 3) -> Optional[Tuple[str, float]]:
    boxes = locate_digit_boxes(mask, expected_digits=expected_digits)
    if boxes is None:
        return None

    digits: List[str] = []
    total_score = 0.0
    for x1, y1, x2, y2 in boxes:
        pad_x = max(1, int(round((x2 - x1) * 0.08)))
        pad_y = max(1, int(round((y2 - y1) * 0.06)))
        x1 = max(0, x1 - pad_x)
        y1 = max(0, y1 - pad_y)
        x2 = min(mask.shape[1], x2 + pad_x)
        y2 = min(mask.shape[0], y2 + pad_y)
        digit, score = decode_digit_mask(mask[y1:y2, x1:x2])
        if digit is None:
            return None
        digits.append(digit)
        total_score += score

    return "".join(digits), total_score / expected_digits


def choose_best_candidate(candidates: Iterable[Optional[str]]) -> Optional[str]:
    values = [candidate for candidate in candidates if candidate]
    if not values:
        return None
    return max(values, key=lambda candidate: (len(candidate) == 3, values.count(candidate), len(candidate)))


def read_digits_from_crop(crop: np.ndarray, color: str, ocr: RapidOCR) -> Optional[str]:
    best_segment: Optional[Tuple[str, float]] = None
    for mask in build_digit_mask_variants(crop, color):
        candidate = read_seven_segment_digits(mask)
        if candidate is None:
            continue
        if best_segment is None or candidate[1] > best_segment[1]:
            best_segment = candidate

    if best_segment is not None:
        return best_segment[0]

    base_mask = threshold_display_digits(crop, color)
    enlarged_color = cv2.resize(crop, None, fx=4, fy=4, interpolation=cv2.INTER_CUBIC)
    gray = cv2.cvtColor(enlarged_color, cv2.COLOR_BGR2GRAY)
    gray = cv2.equalizeHist(gray)
    _, binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    mask_4x = cv2.resize(base_mask, None, fx=4, fy=4, interpolation=cv2.INTER_NEAREST)
    mask_5x = cv2.resize(base_mask, None, fx=5, fy=5, interpolation=cv2.INTER_NEAREST)

    return choose_best_candidate(
        [
            run_ocr_digits(enlarged_color, ocr),
            run_ocr_digits(binary, ocr),
            run_ocr_digits(255 - binary, ocr),
            run_ocr_digits(mask_4x, ocr),
            run_ocr_digits(cv2.erode(mask_4x, np.ones((2, 2), dtype=np.uint8), iterations=1), ocr),
            run_ocr_digits(mask_5x, ocr),
        ]
    )


def extract_display_value(
    display: np.ndarray,
    ratios: Tuple[float, float, float, float],
    color: str,
    ocr: RapidOCR,
) -> Optional[float]:
    crop = crop_by_ratio(display, ratios)
    return parse_fixed_one_decimal(read_digits_from_crop(crop, color, ocr))


def extract_blue_value(display: np.ndarray, ocr: RapidOCR) -> Optional[float]:
    return extract_display_value(display, BLUE_CROP, "blue", ocr)


def extract_red_value(display: np.ndarray, ocr: RapidOCR) -> Optional[float]:
    return extract_display_value(display, RED_CROP, "red", ocr)


def detect_display_roi(first_frame: np.ndarray) -> Tuple[int, int, int, int]:
    rotated = rotate_frame_if_needed(first_frame)
    b, g, r = cv2.split(rotated)
    diff = b.astype(int) - np.maximum(g, r).astype(int)
    mask = (diff > 80).astype(np.uint8) * 255
    num_labels, _, stats, _ = cv2.connectedComponentsWithStats(mask)

    if num_labels <= 1:
        raise RuntimeError("Could not detect the blue display digits in the first video frame")

    component_index = max(
        range(1, num_labels),
        key=lambda index: int(stats[index, cv2.CC_STAT_AREA]),
    )
    x = int(stats[component_index, cv2.CC_STAT_LEFT])
    y = int(stats[component_index, cv2.CC_STAT_TOP])
    w = int(stats[component_index, cv2.CC_STAT_WIDTH])
    h = int(stats[component_index, cv2.CC_STAT_HEIGHT])

    x1 = max(0, x - w)
    y1 = max(0, y - int(round(0.3 * h)))
    x2 = min(rotated.shape[1], x + int(round(1.8 * w)))
    y2 = min(rotated.shape[0], y + int(round(2.5 * h)))
    return x1, y1, x2, y2


def parse_video_start_time(video_path: Path) -> datetime:
    match = TIMESTAMP_RE.search(video_path.stem)
    if match:
        return datetime.strptime("{0}{1}".format(match.group(1), match.group(2)), "%Y%m%d%H%M%S")
    return datetime.fromtimestamp(video_path.stat().st_mtime)


def list_videos(video_dir: Path) -> List[Path]:
    videos = [
        path
        for path in sorted(video_dir.iterdir())
        if path.is_file() and path.suffix.lower() in VIDEO_EXTENSIONS
    ]
    if not videos:
        raise FileNotFoundError("No video files were found in {0}".format(video_dir))
    return videos


def choose_latest_excel(excel_dir: Path) -> Path:
    candidates = []
    for path in excel_dir.glob("*.xlsx"):
        if not path.is_file():
            continue
        if path.name.startswith("~$"):
            continue
        if ".tmp." in path.name:
            continue
        if path.stem.endswith("_video_ocr"):
            continue
        candidates.append(path)

    if not candidates:
        raise FileNotFoundError("No source Excel workbook was found in {0}".format(excel_dir))

    return max(candidates, key=lambda item: item.stat().st_mtime)


def load_video_infos(video_paths: Sequence[Path]) -> List[VideoInfo]:
    infos: List[VideoInfo] = []
    for path in video_paths:
        capture = cv2.VideoCapture(str(path))
        if not capture.isOpened():
            raise RuntimeError("Could not open video: {0}".format(path))
        fps = float(capture.get(cv2.CAP_PROP_FPS) or 0.0)
        frame_count = int(capture.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
        ok, first_frame = capture.read()
        capture.release()
        if not ok or first_frame is None:
            raise RuntimeError("Could not read the first frame from {0}".format(path))
        if fps <= 0:
            raise RuntimeError("Invalid FPS reported for {0}".format(path))

        roi = detect_display_roi(first_frame)
        infos.append(
            VideoInfo(
                path=path,
                start_time=parse_video_start_time(path),
                fps=fps,
                frame_count=frame_count,
                duration_seconds=frame_count / fps if fps > 0 else 0.0,
                roi=roi,
            )
        )
    return sorted(infos, key=lambda item: item.start_time)


def parse_excel_timestamp(value: object) -> Optional[datetime]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def choose_video_for_time(row_time: datetime, videos: Sequence[VideoInfo]) -> Optional[VideoInfo]:
    covering = [video for video in videos if video.start_time <= row_time <= video.end_time]
    if covering:
        return min(covering, key=lambda video: abs((row_time - video.start_time).total_seconds()))
    return None


def ensure_output_headers(worksheet) -> Dict[str, int]:
    headers = [cell.value for cell in worksheet[1]] if worksheet.max_row >= 1 else []
    header_map: Dict[str, int] = {}
    for index, value in enumerate(headers, start=1):
        if value is not None and str(value).strip():
            header_map[str(value)] = index

    next_column = len(headers) + 1
    for header in OUTPUT_HEADERS:
        if header not in header_map:
            worksheet.cell(row=1, column=next_column, value=header)
            header_map[header] = next_column
            next_column += 1
    return header_map


def build_output_path(source_workbook: Path, output_path: Optional[Path]) -> Path:
    if output_path is not None:
        return output_path
    return source_workbook.with_name("{0}_video_ocr{1}".format(source_workbook.stem, source_workbook.suffix))


def process_workbook(
    source_workbook: Path,
    output_workbook: Path,
    videos: Sequence[VideoInfo],
    ocr: RapidOCR,
) -> Tuple[int, int]:
    workbook = load_workbook(source_workbook)
    worksheet = workbook[workbook.sheetnames[0]]
    header_map = ensure_output_headers(worksheet)
    readers = {video.path: VideoReader(video) for video in videos}
    updated_rows = 0
    matched_rows = 0

    try:
        for row_index in range(2, worksheet.max_row + 1):
            row_time = parse_excel_timestamp(worksheet.cell(row=row_index, column=1).value)
            if row_time is None:
                continue

            video = choose_video_for_time(row_time, videos)
            if video is None:
                continue

            seconds_from_start = (row_time - video.start_time).total_seconds()
            frame_index = int(round(seconds_from_start * video.fps))
            voltage, current, frame_time, _ = readers[video.path].read_measurement(frame_index, ocr)
            match_delta_ms = round((frame_time - row_time).total_seconds() * 1000.0, 3)

            worksheet.cell(
                row=row_index,
                column=header_map["video_timestamp"],
                value=frame_time.strftime("%Y-%m-%d %H:%M:%S.%f"),
            )
            worksheet.cell(
                row=row_index,
                column=header_map["video_match_delta_ms"],
                value=match_delta_ms,
            )
            worksheet.cell(
                row=row_index,
                column=header_map["video_voltage"],
                value=voltage,
            )
            worksheet.cell(
                row=row_index,
                column=header_map["video_current"],
                value=current,
            )
            worksheet.cell(
                row=row_index,
                column=header_map["video_file"],
                value=video.path.name,
            )
            updated_rows += 1
            if voltage is not None or current is not None:
                matched_rows += 1

        output_workbook.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_workbook)
        return updated_rows, matched_rows
    finally:
        workbook.close()
        for reader in readers.values():
            reader.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Read voltage/current values from videos and append them to the latest Excel workbook."
    )
    parser.add_argument(
        "--video-dir",
        default="videos",
        help="Folder that contains the videos. Default: videos",
    )
    parser.add_argument(
        "--excel-dir",
        default="logs",
        help="Folder that contains the Excel files. Default: logs",
    )
    parser.add_argument(
        "--output",
        help="Optional output workbook path. Default: <latest_excel>_video_ocr.xlsx",
    )
    return parser


def resolve_input_path(raw_path: str) -> Path:
    path = Path(raw_path)
    if path.is_absolute():
        return path.resolve()
    return (get_app_root() / path).resolve()


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    video_dir = resolve_input_path(args.video_dir)
    excel_dir = resolve_input_path(args.excel_dir)
    source_workbook = choose_latest_excel(excel_dir)
    output_workbook = build_output_path(
        source_workbook,
        resolve_input_path(args.output) if args.output else None,
    )
    video_paths = list_videos(video_dir)
    videos = load_video_infos(video_paths)
    ocr = RapidOCR(use_text_det=False)

    updated_rows, matched_rows = process_workbook(
        source_workbook=source_workbook,
        output_workbook=output_workbook,
        videos=videos,
        ocr=ocr,
    )

    print("Latest source workbook: {0}".format(source_workbook))
    print("Videos processed: {0}".format(len(videos)))
    print("Rows updated: {0}".format(updated_rows))
    print("Rows with OCR values: {0}".format(matched_rows))
    print("Saved annotated workbook: {0}".format(output_workbook))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
